"""Simulateur Excel du budget de diffusion de l'offre « Avocat(e) salarié(e) ».

Usage : python3 generer_budget.py [sortie.xlsx]

Onglets :
- Budget : paramètres, plateformes et tarifs, trois scénarios (quantités modifiables),
  budget du scénario retenu et synthèse comparative ;
- Échéancier : décaissements mois par mois du scénario retenu.

Tous les montants sont calculés par formules : modifier un tarif, une quantité ou la durée
met l'ensemble à jour.
"""

import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

SORTIE = sys.argv[1] if len(sys.argv) > 1 else "Budget_diffusion_offre_avocat.xlsx"

# Charte du cabinet
ARDOISE, ROUGE, ORANGE, VERT, CREME = "304859", "B52026", "FCAF19", "8EC33F", "FAF0C7"
T_VERT, T_ORANGE, T_ROUGE = "EEF6E2", "FFF3DC", "F7E6E6"
SAISIE, CALCUL, FILET = "FFF2CC", "F2F2F2", "BFBFBF"
POLICE = "Arial"
EURO = '#,##0 "€";-#,##0 "€";"–"'
fin = Side(style="thin", color=FILET)
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)


def f(**kw):
    kw.setdefault("name", POLICE)
    kw.setdefault("size", 10)
    return Font(**kw)


def fond(c):
    return PatternFill("solid", start_color=c, end_color=c)


def bandeau(ws, plage, texte, couleur=ARDOISE, texte_couleur="FFFFFF"):
    debut, fin_ = plage.split(":")
    if debut != fin_:
        ws.merge_cells(plage)
    c = ws[plage.split(":")[0]]
    c.value = texte
    c.font = f(bold=True, color=texte_couleur)
    c.fill = fond(couleur)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def saisie(c, fmt=None):
    c.fill = fond(SAISIE)
    c.font = f(color="0000FF")
    c.border = BORD
    if fmt:
        c.number_format = fmt


def calcul(c, fmt=None, gras=False):
    c.fill = fond(CALCUL)
    c.font = f(bold=gras)
    c.border = BORD
    if fmt:
        c.number_format = fmt


wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.sheet_view.showGridLines = False
for col, l in {"A": 2, "B": 31, "C": 27, "D": 11, "E": 13, "F": 11, "G": 10, "H": 12, "I": 10, "J": 12,
               "K": 10, "L": 12, "M": 44, "N": 3, "O": 11, "P": 12, "Q": 12, "R": 12, "S": 12, "T": 12, "U": 12}.items():
    ws.column_dimensions[col].width = l

ws["B1"] = "Budget de diffusion de l'offre – Avocat(e) salarié(e)"
ws["B1"].font = f(size=16, bold=True, color=ARDOISE)
ws["B2"] = ("Victimes & Préjudices Avocats · estimation du budget mensuel par plateforme et par scénario. "
            "Cellules jaunes = saisie ; cellules grises = calcul automatique.")
ws["B2"].font = f(size=9, italic=True, color="595959")

# --------------------------------------------------------------------------
# Paramètres
# --------------------------------------------------------------------------
bandeau(ws, "B4:D4", "PARAMÈTRES")
PARAMS = [
    (5, "Début de la diffusion", dt.date(2026, 10, 1), "DD/MM/YYYY",
     "Premier mois de l'échéancier."),
    (6, "Durée de la campagne (mois)", 2, "0",
     "Durée prévue de la recherche. Le forfait Village de la Justice couvre 2 mois : au-delà, il est renouvelé."),
    (7, "Tarifs saisis en", "HT", None, "HT ou TTC : sert à calculer l'autre montant dans le budget retenu."),
    (8, "Taux de TVA", 0.2, "0%", None),
    (9, "Scénario retenu", "Équilibré", None, "Essentiel, Équilibré ou Intensif : pilote le budget retenu et l'échéancier."),
]
for r, libelle, valeur, fmt, note in PARAMS:
    ws[f"B{r}"] = libelle
    ws[f"B{r}"].font = f()
    ws[f"B{r}"].border = BORD
    ws.merge_cells(f"C{r}:D{r}")
    ws[f"C{r}"] = valeur
    saisie(ws[f"C{r}"], fmt)
    ws[f"D{r}"].border = BORD
    ws[f"C{r}"].alignment = Alignment(horizontal="center")
    if note:
        ws[f"B{r}"].comment = Comment(note, "Simulateur")
ws["C9"].font = f(color="0000FF", bold=True)

dv_ht = DataValidation(type="list", formula1='"HT,TTC"', allow_blank=False)
dv_scen = DataValidation(type="list", formula1="=$B$29:$B$31", allow_blank=False)
dv_unite = DataValidation(type="list", formula1='"Gratuit,Forfait,Par jour,Par semaine,Par mois"', allow_blank=True)
dv_duree = DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                          showErrorMessage=True, error="Durée entre 1 et 12 mois.")
for dv in (dv_ht, dv_scen, dv_unite, dv_duree):
    ws.add_data_validation(dv)
dv_ht.add("C7")
dv_scen.add("C9")
dv_duree.add("C6")

# --------------------------------------------------------------------------
# Plateformes et scénarios
# --------------------------------------------------------------------------
PREM, DERN = 13, 24  # lignes des plateformes (10 prévues + 2 libres)
bandeau(ws, "B11:F11", "PLATEFORMES ET TARIFS")
bandeau(ws, "G11:H11", "ESSENTIEL", VERT, "FFFFFF")
bandeau(ws, "I11:J11", "ÉQUILIBRÉ", ORANGE, ARDOISE)
bandeau(ws, "K11:L11", "INTENSIF", ROUGE, "FFFFFF")
bandeau(ws, "M11:M11", "COMMENTAIRES")
ENTETES = ["Plateforme", "Formule", "Tarif (€)", "Facturation", "Durée du forfait (mois)",
           "Qté / mois", "Coût / mois", "Qté / mois", "Coût / mois", "Qté / mois", "Coût / mois", ""]
for col, texte in zip("BCDEFGHIJKLM", ENTETES):
    c = ws[f"{col}12"]
    c.value = texte
    c.font = f(bold=True, size=9)
    c.fill = fond("D9D9D9")
    c.border = BORD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[12].height = 30
for col in "GIK":
    ws[f"{col}12"].comment = Comment(
        "Quantité par mois selon la facturation :\n– gratuit ou forfait : 1 = plateforme retenue, 0 = non ;\n"
        "– par jour : nombre de jours de boost dans le mois ;\n– par semaine : nombre de semaines sponsorisées dans le mois ;\n"
        "– par mois : 1.", "Simulateur")

# (plateforme, formule, tarif, facturation, durée forfait, qté Essentiel, Équilibré, Intensif, commentaire)
PLATEFORMES = [
    ("Village de la Justice", "Annonce Premium", 290, "Forfait", 2, 1, 1, 1,
     "290 € pour 2 mois, payés à la publication. Site de référence des professions du droit."),
    ("RecrutAvocat", "Annonce", 0, "Gratuit", None, 1, 1, 1, "Site spécialisé dans le recrutement d'avocats."),
    ("LinkedIn", "Offre d'emploi", 0, "Gratuit", None, 1, 1, 1, "Publication de l'offre gratuite."),
    ("LinkedIn", "Boost de l'offre", 10, "Par jour", None, 0, 15, 30,
     "Budget journalier maximal : la dépense réelle peut être inférieure."),
    ("Indeed", "Annonce sponsorisée", 100, "Par semaine", None, 0, 2, 4, "Facturé à la semaine sponsorisée."),
    ("France Travail", "Offre d'emploi", 0, "Gratuit", None, 1, 1, 1, ""),
    ("Apec", "Offre cadre", 0, "Gratuit", None, 1, 1, 1, "Poste cadre : cible les avocats salariés en poste ou en recherche."),
    ("CNB", "Offre d'emploi", 0, "Gratuit", None, 1, 1, 1, "Conseil national des barreaux."),
    ("Ordre des avocats de Grenoble", "Offre d'emploi", 0, "Gratuit", None, 1, 1, 1, "Diffusion auprès du barreau local."),
    ("Meta (Facebook, Instagram)", "Boost de publication", 10, "Par jour", None, 0, 10, 30,
     "Budget journalier maximal : la dépense réelle peut être inférieure."),
    ("Autre plateforme", None, None, None, None, None, None, None, "Ligne libre."),
    ("Autre plateforme", None, None, None, None, None, None, None, "Ligne libre."),
]


def cout_mensuel(r, q):
    """Coût mensuel lissé : forfait réparti sur sa durée, sinon tarif × quantité."""
    return (f'=IF(OR(${q}{r}="",N(${q}{r})=0,$E{r}="",$E{r}="Gratuit"),0,'
            f'IF($E{r}="Forfait",N($D{r})/MAX(1,N($F{r})),N($D{r})*${q}{r}))')


def premier_mois(r, q):
    """Décaissement du premier mois : forfait payé en entier."""
    return (f'=IF(OR(${q}{r}="",N(${q}{r})=0,$E{r}="",$E{r}="Gratuit"),0,'
            f'IF($E{r}="Forfait",N($D{r}),N($D{r})*${q}{r}))')


def campagne(r, q):
    """Coût sur la durée de la campagne : forfaits renouvelés au besoin."""
    return (f'=IF(OR(${q}{r}="",N(${q}{r})=0,$E{r}="",$E{r}="Gratuit"),0,'
            f'IF($E{r}="Forfait",N($D{r})*ROUNDUP($C$6/MAX(1,N($F{r})),0),N($D{r})*${q}{r}*$C$6))')


# Détail des calculs (colonnes O à U)
bandeau(ws, "O11:U11", "DÉTAIL DES CALCULS (ne pas modifier)", "7F7F7F")
for col, texte in zip("OPQRSTU", ["Qté du scénario retenu", "1er mois Essentiel", "Campagne Essentiel",
                                  "1er mois Équilibré", "Campagne Équilibré", "1er mois Intensif", "Campagne Intensif"]):
    c = ws[f"{col}12"]
    c.value = texte
    c.font = f(bold=True, size=8, color="595959")
    c.fill = fond("EDEDED")
    c.border = BORD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for i, (nom, formule, tarif, unite, duree, qe, qq, qi, com) in enumerate(PLATEFORMES):
    r = PREM + i
    for col, val in zip("BCDEF", [nom, formule, tarif, unite, duree]):
        ws[f"{col}{r}"] = val
        saisie(ws[f"{col}{r}"])
    ws[f"D{r}"].number_format = EURO
    ws[f"F{r}"].number_format = "0"
    for col in "DEF":
        ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
    for q, val in zip("GIK", [qe, qq, qi]):
        ws[f"{q}{r}"] = val
        saisie(ws[f"{q}{r}"], "0")
        ws[f"{q}{r}"].alignment = Alignment(horizontal="center")
    for q, c in zip("GIK", "HJL"):
        ws[f"{c}{r}"] = cout_mensuel(r, q)
        calcul(ws[f"{c}{r}"], EURO)
    ws[f"M{r}"] = com
    ws[f"M{r}"].font = f(size=9, color="595959")
    ws[f"M{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    ws[f"M{r}"].border = BORD
    # détail
    ws[f"O{r}"] = f"=CHOOSE($C$10,G{r},I{r},K{r})"
    ws[f"P{r}"], ws[f"Q{r}"] = premier_mois(r, "G"), campagne(r, "G")
    ws[f"R{r}"], ws[f"S{r}"] = premier_mois(r, "I"), campagne(r, "I")
    ws[f"T{r}"], ws[f"U{r}"] = premier_mois(r, "K"), campagne(r, "K")
    for col in "OPQRSTU":
        calcul(ws[f"{col}{r}"], "0" if col == "O" else EURO)
        ws[f"{col}{r}"].font = f(size=9, color="595959")
    ws.row_dimensions[r].height = 26
dv_unite.add(f"E{PREM}:E{DERN}")

# Totaux
RT = DERN + 1
ws[f"B{RT}"] = "Total mensuel (forfait lissé sur sa durée)"
ws.merge_cells(f"B{RT}:F{RT}")
ws[f"B{RT}"].font = f(bold=True)
ws[f"B{RT}"].alignment = Alignment(horizontal="right")
for col in "GHIJKL":
    ws[f"{col}{RT}"].border = BORD
for q, c, teinte in zip("GIK", "HJL", [T_VERT, T_ORANGE, T_ROUGE]):
    ws[f"{q}{RT}"] = f'=COUNTIFS({q}{PREM}:{q}{DERN},">0",$E{PREM}:$E{DERN},"<>Gratuit")'
    ws[f"{q}{RT}"].number_format = '0" payante(s)"'
    ws[f"{q}{RT}"].font = f(size=8, color="595959")
    ws[f"{q}{RT}"].fill = fond(teinte)
    ws[f"{c}{RT}"] = f"=SUM({c}{PREM}:{c}{DERN})"
    ws[f"{c}{RT}"].number_format = EURO
    ws[f"{c}{RT}"].font = f(bold=True)
    ws[f"{c}{RT}"].fill = fond(teinte)
for col in "PQRSTU":
    ws[f"{col}{RT}"] = f"=SUM({col}{PREM}:{col}{DERN})"
    calcul(ws[f"{col}{RT}"], EURO, gras=True)
ws["O25"].border = BORD

# Index du scénario retenu (cellule technique)
ws["B10"] = "Numéro du scénario retenu"
ws["B10"].font = f(size=8, color="A6A6A6")
ws["C10"] = "=MATCH($C$9,$B$29:$B$31,0)"
ws["C10"].font = f(size=8, color="A6A6A6")
ws["C10"].alignment = Alignment(horizontal="center")
ws.row_dimensions[10].hidden = True  # ligne technique

# --------------------------------------------------------------------------
# Budget du scénario retenu (indicateurs)
# --------------------------------------------------------------------------
bandeau(ws, "F4:L4", "BUDGET DU SCÉNARIO RETENU")
ws["F5"] = '="Scénario "&$C$9&" · campagne de "&$C$6&" mois"'
ws.merge_cells("F5:J5")
ws["F5"].font = f(size=9, italic=True, color="595959")
for col, texte in (("K", "HT"), ("L", "TTC")):
    ws[f"{col}5"] = texte
    ws[f"{col}5"].font = f(bold=True, size=9)
    ws[f"{col}5"].alignment = Alignment(horizontal="center")
    ws[f"{col}5"].fill = fond("D9D9D9")
    ws[f"{col}5"].border = BORD

HT = lambda x: f'=IF($C$7="HT",{x},{x}/(1+$C$8))'
TTC = lambda x: f'=IF($C$7="HT",{x}*(1+$C$8),{x})'
TOTAL_RETENU = f"CHOOSE($C$10,$Q${RT},$S${RT},$U${RT})"
PREMIER_RETENU = f"CHOOSE($C$10,$P${RT},$R${RT},$T${RT})"
MOYEN_RETENU = f"{TOTAL_RETENU}/$C$6"
INDICATEURS = [
    (6, "Budget mensuel moyen à allouer", MOYEN_RETENU),
    (7, "Décaissement du 1er mois (forfait payé d'avance)", PREMIER_RETENU),
    (8, "Budget total de la campagne", TOTAL_RETENU),
]
for r, libelle, expr in INDICATEURS:
    ws.merge_cells(f"F{r}:J{r}")
    ws[f"F{r}"] = libelle
    ws[f"F{r}"].font = f(bold=(r == 6))
    ws[f"F{r}"].border = BORD
    ws[f"K{r}"] = HT(expr)
    ws[f"L{r}"] = TTC(expr)
    for col in "KL":
        calcul(ws[f"{col}{r}"], EURO, gras=True)
        ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
ws.merge_cells("F9:J9")
ws["F9"] = "Plateformes utilisées (dont payantes)"
ws["F9"].font = f()
ws["F9"].border = BORD
ws["K9"] = f'=COUNTIF($O${PREM}:$O${DERN},">0")&" ("&CHOOSE($C$10,$G${RT},$I${RT},$K${RT})&")"'
ws.merge_cells("K9:L9")
calcul(ws["K9"])
ws["K9"].alignment = Alignment(horizontal="center")
# indicateur principal mis en valeur
for col in "KL":
    ws[f"{col}6"].fill = fond(CREME)
    ws[f"{col}6"].font = f(bold=True, size=12, color=ROUGE)

# --------------------------------------------------------------------------
# Synthèse des scénarios
# --------------------------------------------------------------------------
bandeau(ws, "B27:H27", "SYNTHÈSE DES SCÉNARIOS")
for col, texte in zip("BCDEFGH", ["Scénario", "Contenu", "Mensuel moyen", "1er mois", "Total campagne",
                                  "Total TTC", "Plateformes payantes"]):
    c = ws[f"{col}28"]
    c.value = texte
    c.font = f(bold=True, size=9)
    c.fill = fond("D9D9D9")
    c.border = BORD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
SCENARIOS = [
    (29, "Essentiel", "Plateformes gratuites + Village de la Justice", "Q", "P", "G", T_VERT),
    (30, "Équilibré", "Essentiel + boosts LinkedIn (15 j) et Meta (10 j) + Indeed 2 semaines par mois", "S", "R", "I", T_ORANGE),
    (31, "Intensif", "Essentiel + boosts LinkedIn et Meta tous les jours + Indeed chaque semaine", "U", "T", "K", T_ROUGE),
]
for r, nom, contenu, camp, prem, qcol, teinte in SCENARIOS:
    ws[f"B{r}"] = nom
    ws[f"B{r}"].font = f(bold=True)
    ws[f"B{r}"].fill = fond(teinte)
    ws[f"C{r}"] = contenu
    ws[f"C{r}"].font = f(size=8, color="595959")
    ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    ws[f"D{r}"] = f"=${camp}${RT}/$C$6"
    ws[f"E{r}"] = f"=${prem}${RT}"
    ws[f"F{r}"] = f"=${camp}${RT}"
    ws[f"G{r}"] = f'=IF($C$7="HT",F{r}*(1+$C$8),F{r})'
    ws[f"H{r}"] = f"=${qcol}${RT}"
    for col in "DEFG":
        calcul(ws[f"{col}{r}"], EURO)
    calcul(ws[f"H{r}"], "0")
    for col in "BCDEFGH":
        ws[f"{col}{r}"].border = BORD
        if col != "C":
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 30
ws["C29"].comment = Comment("Les quantités de chaque scénario sont modifiables dans le tableau des plateformes.", "Simulateur")
ws["B32"] = ("Montants exprimés selon la base choisie dans « Tarifs saisis en » (HT ou TTC). "
             "Mensuel moyen = total de la campagne ÷ durée.")
ws["B32"].font = f(size=8, italic=True, color="595959")

# Mise en évidence du scénario retenu
for r, nom, *_ in SCENARIOS:
    ws.conditional_formatting.add(f"B{r}:H{r}", FormulaRule(formula=[f'$C$9="{nom}"'],
                                  font=Font(name=POLICE, bold=True, color=ROUGE)))
for plage, nom in (("G11:H12", "Essentiel"), ("I11:J12", "Équilibré"), ("K11:L12", "Intensif")):
    ws.conditional_formatting.add(plage, FormulaRule(formula=[f'$C$9="{nom}"'],
                                  border=Border(bottom=Side(style="thick", color=ARDOISE))))

# --------------------------------------------------------------------------
# Hypothèses
# --------------------------------------------------------------------------
bandeau(ws, "B34:M34", "HYPOTHÈSES ET MODE D'EMPLOI")
NOTES = [
    "Tarifs communiqués par le cabinet (septembre 2026). Vérifiez s'ils sont HT ou TTC et ajustez « Tarifs saisis en ».",
    "Village de la Justice : forfait de 290 € pour 2 mois, payé à la publication. Il est renouvelé si la campagne dure plus de 2 mois.",
    "Boosts LinkedIn et Meta : 10 € par jour correspondent à un budget maximal ; la dépense réelle peut être inférieure. "
    "Saisissez le nombre de jours de boost prévus chaque mois.",
    "Indeed : 100 € par semaine sponsorisée ; saisissez le nombre de semaines par mois (4 au maximum pour un mois plein).",
    "Plateformes gratuites : quantité 1 = diffusion prévue, 0 = pas de diffusion. Elles n'entrent pas dans le budget mais dans le "
    "décompte des plateformes utilisées.",
    "Pour tester une autre hypothèse, modifiez les quantités d'un scénario ; pour ajouter un support, utilisez les lignes "
    "« Autre plateforme ».",
]
for i, note in enumerate(NOTES):
    r = 35 + i
    ws.merge_cells(f"B{r}:M{r}")
    ws[f"B{r}"] = "• " + note
    ws[f"B{r}"].font = f(size=9)
    ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26

ws.freeze_panes = "A4"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
ws.print_area = f"A1:M{35 + len(NOTES)}"

# --------------------------------------------------------------------------
# Échéancier du scénario retenu
# --------------------------------------------------------------------------
we = wb.create_sheet("Échéancier")
we.sheet_view.showGridLines = False
we.column_dimensions["A"].width = 2
we.column_dimensions["B"].width = 34
for k in range(12):
    we.column_dimensions[chr(ord("C") + k)].width = 11
we.column_dimensions["O"].width = 13
we["B1"] = "Échéancier mensuel des dépenses"
we["B1"].font = f(size=16, bold=True, color=ARDOISE)
we["B2"] = '="Scénario "&Budget!$C$9&" · campagne de "&Budget!$C$6&" mois à partir de "&TEXT(MONTH(Budget!$C$5),"00")&"/"&YEAR(Budget!$C$5)'
we["B2"].font = f(size=10, italic=True, color="595959")
we["B3"] = "Décaissements prévus par mois (forfaits payés au début de leur période). Aucune saisie dans cet onglet."
we["B3"].font = f(size=8, italic=True, color="595959")

bandeau(we, "B5:O5", "DÉPENSES PAR PLATEFORME ET PAR MOIS")
we["B6"] = "Mois de campagne"
we["B7"] = "Plateforme"
for k in range(12):
    col = chr(ord("C") + k)
    we[f"{col}6"] = k + 1
    we[f"{col}7"] = f"=EDATE(Budget!$C$5,{k})"
    we[f"{col}7"].number_format = "MMM YYYY"
    for rr in (6, 7):
        we[f"{col}{rr}"].alignment = Alignment(horizontal="center")
        we[f"{col}{rr}"].border = BORD
    we[f"{col}6"].font = f(size=8, color="595959")
    we[f"{col}7"].font = f(bold=True, size=9)
    we[f"{col}7"].fill = fond("D9D9D9")
we["O7"] = "Total"
we["O7"].font = f(bold=True, size=9)
we["O7"].fill = fond("D9D9D9")
we["O7"].alignment = Alignment(horizontal="center")
for c in ("B6", "B7", "O6", "O7"):
    we[c].border = BORD
we["B6"].font = f(size=8, color="595959")
we["B7"].font = f(bold=True, size=9)
we["B7"].fill = fond("D9D9D9")

for i in range(DERN - PREM + 1):
    r = 8 + i
    rb = PREM + i
    we[f"B{r}"] = (f'=IF(Budget!$B${rb}="","",Budget!$B${rb}&IF(Budget!$C${rb}="",""," – "&Budget!$C${rb}))')
    we[f"B{r}"].font = f(size=9)
    we[f"B{r}"].border = BORD
    for k in range(12):
        col = chr(ord("C") + k)
        we[f"{col}{r}"] = (
            f'=IF(OR({col}$6>Budget!$C$6,N(Budget!$O${rb})=0,Budget!$E${rb}="",Budget!$E${rb}="Gratuit"),0,'
            f'IF(Budget!$E${rb}="Forfait",IF(MOD({col}$6-1,MAX(1,N(Budget!$F${rb})))=0,N(Budget!$D${rb}),0),'
            f'N(Budget!$D${rb})*Budget!$O${rb}))')
        calcul(we[f"{col}{r}"], EURO)
        we[f"{col}{r}"].font = f(size=9)
    we[f"O{r}"] = f"=SUM(C{r}:N{r})"
    calcul(we[f"O{r}"], EURO, gras=True)

rt = 8 + (DERN - PREM + 1)
we[f"B{rt}"] = '="Total du mois ("&Budget!$C$7&")"'
we[f"B{rt + 1}"] = '=IF(Budget!$C$7="HT","Total du mois TTC","Total du mois HT")'
we[f"B{rt + 2}"] = "Cumul depuis le début"
for k in range(13):
    col = chr(ord("C") + k)
    we[f"{col}{rt}"] = f"=SUM({col}8:{col}{rt - 1})"
    we[f"{col}{rt + 1}"] = f'=IF(Budget!$C$7="HT",{col}{rt}*(1+Budget!$C$8),{col}{rt}/(1+Budget!$C$8))'
    if col != "O":
        we[f"{col}{rt + 2}"] = f"=SUM($C${rt}:{col}{rt})"
for rr, teinte in ((rt, CREME), (rt + 1, "FFFFFF"), (rt + 2, "FFFFFF")):
    we[f"B{rr}"].font = f(bold=(rr == rt), size=9)
    we[f"B{rr}"].border = BORD
    for k in range(13):
        col = chr(ord("C") + k)
        c = we[f"{col}{rr}"]
        if c.value is None:
            continue
        c.number_format = EURO
        c.border = BORD
        c.fill = fond(teinte)
        c.font = f(bold=(rr == rt), size=9, color=ROUGE if rr == rt else "000000")
# Colonnes hors campagne grisées
we.conditional_formatting.add("C6:N" + str(rt + 2), FormulaRule(formula=["C$6>Budget!$C$6"],
                              font=Font(name=POLICE, color="BFBFBF")))
we.freeze_panes = "C8"
we.page_setup.orientation = "landscape"
we.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
we.page_setup.fitToHeight = 0

wb.active = 0
wb.save(SORTIE)
print(f"Simulateur généré : {SORTIE}")
