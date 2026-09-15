# -*- coding: utf-8 -*-
"""Contrôles de cohérence du classeur, du module VBA et des modèles Word."""
import re, sys, glob, zipfile, logging, warnings
from xml.etree import ElementTree as ET
warnings.filterwarnings("ignore"); logging.disable(logging.CRITICAL)
from openpyxl import load_workbook

XLSX = "Gestion_Candidatures_VP.xlsx"
BAS  = "VP_Candidatures.bas"
DOCX = sorted(glob.glob("modeles-word/*.docx"))
ERREURS, OK = [], []

wb  = load_workbook(XLSX)
vba = open(BAS, encoding="cp1252").read()

# 1. feuilles referencees par le VBA
feuilles = re.findall(r'Const SH_\w+\s+As String = "([^"]+)"', vba)
manq = [f for f in feuilles if f not in wb.sheetnames]
(ERREURS if manq else OK).append("Feuilles VBA -> classeur : %s" % (manq or "%d/%d" % (len(feuilles), len(feuilles))))

# 2. en-tetes de colonnes manipules par le VBA
hdr = [c.value for c in wb["Candidatures"][4] if c.value]
used = set(re.findall(r'(?:Txt_|Val_)\(\s*[\w()]+\s*,\s*"([^"]+)"', vba)) | \
       set(re.findall(r'Ecrire\s+[\w()]+\s*,\s*"([^"]+)"', vba))
manq = sorted(h for h in used if h not in hdr)
(ERREURS if manq else OK).append("En-têtes VBA -> Candidatures : %s" % (manq or "%d/%d" % (len(used), len(used))))

# 3. plages nommees
noms = set(re.findall(r'Names\("(\w+)"\)', vba)) | set(re.findall(r'Param(?:Txt|Num)?\("(\w+)"', vba))
for bloc in re.findall(r'noms = Array\((.*?)\)\n', vba, re.S):
    noms |= set(re.findall(r'"(\w+)"', bloc))
manq = sorted(n for n in noms if n not in wb.defined_names)
(ERREURS if manq else OK).append("Plages nommées VBA -> classeur : %s" % (manq or "%d/%d" % (len(noms), len(noms))))

# 4. macros des boutons
publics = set(re.findall(r'^Public Sub (\w+)\(\)', vba, re.M))
boutons = {m for m in re.findall(r'Array\("(\w+)",\s*"', vba) if m[0].isupper() and m in publics or
           (m[0].isupper() and m not in ("Civilites",))}
manq = sorted(b for b in boutons if b not in publics)
(ERREURS if manq else OK).append("Boutons -> macros publiques : %s" % (manq or "%d boutons" % len(boutons)))

# 5. equilibrage des procedures / guillemets
n = [len(re.findall(r'^\s*(?:Public |Private )?Sub \w+', vba, re.M)), len(re.findall(r'^\s*End Sub\s*$', vba, re.M)),
     len(re.findall(r'^\s*(?:Public |Private )?Function \w+', vba, re.M)), len(re.findall(r'^\s*End Function\s*$', vba, re.M))]
(ERREURS if (n[0] != n[1] or n[2] != n[3]) else OK).append("Structure VBA : %d Sub / %d End Sub, %d Function / %d End Function" % tuple(n))
impairs = [i for i, l in enumerate(vba.split("\n"), 1) if not l.strip().startswith("'") and l.count('"') % 2]
(ERREURS if impairs else OK).append("Guillemets VBA : %s" % (("lignes " + str(impairs[:5])) if impairs else "tous équilibrés"))

# 6. balises des modeles (e-mails + Word) couvertes par le VBA
gerees = set(re.findall(r'"(\{\{[A-Z_]+\}\})"', vba)) | {"{{SIGNATURE_BLOC}}"}
wm = wb["Modèles e-mails"]
for r in range(4, 12):
    cle = wm.cell(row=r, column=1).value
    if not cle: continue
    txt = (wm.cell(row=r, column=3).value or "") + " " + (wm.cell(row=r, column=4).value or "")
    manq = sorted(set(re.findall(r'\{\{[A-Z_]+\}\}', txt)) - gerees)
    (ERREURS if manq else OK).append("Modèle e-mail %-18s : %s" % (cle, manq or "balises OK"))
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
for f in DOCX:
    z = zipfile.ZipFile(f)
    souci = []
    if z.testzip() is not None: souci.append("archive corrompue")
    for req in ("[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels"):
        if req not in z.namelist(): souci.append("partie absente " + req)
    xmls = ""
    for nm in z.namelist():
        if nm.endswith((".xml", ".rels")):
            try: ET.fromstring(z.read(nm))
            except Exception: souci.append("XML invalide " + nm)
            if nm.startswith("word/"): xmls += z.read(nm).decode("utf-8")
    manq = sorted(set(re.findall(r'\{\{[A-Z_]+\}\}', xmls)) - gerees)
    (ERREURS if (souci or manq) else OK).append("Modèle Word %-28s : %s" % (f.split("/")[-1], souci + manq or "valide, balises OK"))

# 7. formules du classeur
try:
    import formulas
    xl = formulas.ExcelModel().loads(XLSX).finish()
    sol = xl.calculate()
    codes = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!", "#CYCLE!", "#ERROR!")
    bad, tot = [], 0
    for k, v in sol.items():
        try: s = str(v.value[0, 0])
        except Exception:
            try: s = str(v.value)
            except Exception: continue
        tot += 1
        if any(c in s for c in codes): bad.append(k.split("]")[-1])
    (ERREURS if bad else OK).append("Formules : %d cellules évaluées, %d erreur(s) %s" % (tot, len(bad), bad[:5] if bad else ""))
except ImportError:
    OK.append("Formules : bibliothèque « formulas » absente, contrôle ignoré")

print("=" * 78)
for o in OK: print("  OK    ", o)
for e in ERREURS: print("  ERREUR", e)
print("=" * 78)
print(("ECHEC : %d probleme(s)" % len(ERREURS)) if ERREURS else "TOUS LES CONTROLES SONT PASSES")
sys.exit(1 if ERREURS else 0)
