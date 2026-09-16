# -*- coding: utf-8 -*-
"""
Générateur du classeur « Gestion des candidatures » — Cabinet Victimes & Préjudices.
Produit Gestion_Candidatures_VP.xlsx (8 onglets, formules, listes déroulantes,
mises en forme conditionnelles, tableau de bord) prêt à recevoir le module VBA
VP_Candidatures.bas (automatisations Outlook + Word).
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import FormulaRule, CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

# ---------------------------------------------------------------- charte
NAVY   = "12355B"   # bandeaux, en-têtes
NAVY_D = "0B2239"
BLUE_L = "E8EEF5"   # fonds clairs
BLUE_2 = "F4F7FB"
GOLD   = "B8912A"   # accent
GREY   = "5A6B7B"
INK    = "1F2933"
WHITE  = "FFFFFF"
OK_G   = "1E8449"
WARN_O = "C87F0A"
BAD_R  = "B03A2E"
FILL_IN = "FFF8E1"  # cellules à remplir par le cabinet
FILL_CALC = "EFF3F7"  # colonnes calculées automatiquement

MAXROW = 5000       # portée des formules du tableau de bord
NDATA  = 30         # lignes préformatées dans l'onglet Candidatures

def F(sz=10, b=False, color=INK, it=False, u=None):
    return Font(name="Arial", size=sz, bold=b, color=color, italic=it, underline=u)

def P(c):
    return PatternFill("solid", fgColor=c)

THIN = Side(style="thin", color="C5D0DB")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def bandeau(ws, last_col, titre, sous_titre=""):
    """Bandeau de titre sur 2 lignes."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(row=1, column=1, value=titre)
    c.font = F(15, True, WHITE); c.fill = P(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    if sous_titre:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        c2 = ws.cell(row=2, column=1, value=sous_titre)
        c2.font = F(9, False, WHITE); c2.fill = P(NAVY_D)
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 18

def lien(ws, coord, sheet, libelle, cell="A1"):
    c = ws[coord]
    c.value = libelle
    c.hyperlink = Hyperlink(ref=coord, location="'%s'!%s" % (sheet, cell))
    c.font = F(10, True, NAVY, u="single")
    return c

wb = Workbook()
wb.calculation.fullCalcOnLoad = True

# ================================================================ PARAMÈTRES
ws = wb.active
ws.title = "Paramètres"
bandeau(ws, 14, "PARAMÈTRES DU CABINET",
        "Toutes les cellules sur fond crème sont à personnaliser. Les listes ci-dessous alimentent les menus déroulants de l'onglet Candidatures.")
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 46
ws.column_dimensions["D"].width = 3
ws.column_dimensions["E"].width = 58

def section(ws, row, titre, col=2, span=2):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span)
    c = ws.cell(row=row, column=col, value=titre)
    c.font = F(11, True, WHITE); c.fill = P(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = 22

def param(ws, row, label, value, name=None, note="", fmt=None):
    ws.cell(row=row, column=2, value=label).font = F(10, True)
    c = ws.cell(row=row, column=3, value=value)
    c.fill = P(FILL_IN); c.border = BOX; c.font = F(10)
    c.alignment = Alignment(vertical="center", wrap_text=False, indent=1)
    if fmt:
        c.number_format = fmt
    if note:
        n = ws.cell(row=row, column=5, value=note)
        n.font = F(9, False, GREY, it=True)
    if name:
        wb.defined_names.add(DefinedName(name, attr_text="'Paramètres'!$C$%d" % row))
    ws.row_dimensions[row].height = 17
    return c

section(ws, 4, "IDENTITÉ DU CABINET")
param(ws,  5, "Nom du cabinet",        "Cabinet Victimes & Préjudices", "NomCabinet")
param(ws,  6, "Adresse (1 ligne)",     "— à compléter —", "AdresseCabinet")
param(ws,  7, "Téléphone",             "— à compléter —", "TelCabinet")
param(ws,  8, "E-mail recrutement",    "recrutement@victimesetprejudices.fr", "EmailRH",
      "Adresse affichée dans les courriers et les modèles d'e-mails.")
param(ws,  9, "Site internet",         "https://www.victimesetprejudices.fr", "SiteWeb")
param(ws, 10, "Signataire des courriers", "— à compléter —", "Signataire")
param(ws, 11, "Fonction du signataire", "Avocat associé", "FonctionSignataire")
param(ws, 13, "Ville (lieu de signature)", "— à compléter —", "VilleCabinet",
      "Figure en tête des courriers : « Lyon, le 15/09/2026 ».")
c = param(ws, 12, "Signature e-mail", "Cordialement,", "SignatureEmail",
          "Laissez vide pour conserver la signature Outlook par défaut (recommandé).")

section(ws, 15, "AUTOMATISATIONS (Outlook / Word)")
param(ws, 16, "Mode d'envoi des e-mails", "Afficher avant envoi", "ModeEnvoi",
      "« Afficher avant envoi » (recommandé) ou « Envoyer directement ».")
param(ws, 17, "Dossier des candidats", r"C:\Cabinet\Recrutement\Candidats", "DossierCandidats",
      "Un sous-dossier par candidat y est créé automatiquement (CV, LM, pièces).")
param(ws, 18, "Dossier des modèles Word", r"C:\Cabinet\Recrutement\Modeles", "DossierModeles",
      "Déposez-y les fichiers .docx fournis avec ce classeur.")
param(ws, 19, "Dossier Outlook à importer", "Candidatures", "DossierOutlook",
      "Nom du dossier de la boîte de réception scanné par « Importer depuis Outlook ».")
param(ws, 20, "Générer aussi un PDF", "Oui", "GenererPDF",
      "Chaque document Word produit est également enregistré en PDF.")

section(ws, 22, "DÉLAIS ET CONFORMITÉ")
param(ws, 23, "Délai d'accusé de réception (jours)", 7, "SeuilAccuse",
      "Hypothèse cabinet, modifiable : au-delà, la candidature passe en « À traiter ».", "0")
param(ws, 24, "Délai de relance (jours)", 21, "SeuilRelance",
      "Hypothèse cabinet, modifiable : utilisé par le bouton « Relances ».", "0")
param(ws, 25, "Conservation des données (années)", 2, "DureeConservation",
      "Recommandation CNIL : 2 ans à compter du dernier contact avec le candidat.", "0")
c = param(ws, 26, "Date du jour (automatique)", "=TODAY()", "Aujourdhui",
          "Cellule calculée — ne pas modifier. Sert à tous les indicateurs de délai.", "dd/mm/yyyy")
c.fill = P(FILL_CALC)

# ---- listes déroulantes -------------------------------------------------
LISTES = {
    "Civilites":   ["M.", "Mme"],
    "Postes":      ["Avocat collaborateur", "Élève-avocat (stage PPI)", "Juriste dommage corporel",
                    "Stagiaire M1 / M2", "Assistant(e) juridique", "Secrétaire juridique",
                    "Alternant(e)", "Chargé(e) de recouvrement", "Stage découverte", "Autre"],
    "Sources":     ["Candidature spontanée", "Site du cabinet", "LinkedIn", "Village de la Justice",
                    "Indeed", "École / Université", "Cooptation", "Barreau / CNB", "Pôle emploi", "Autre"],
    "StatutListe": ["1-Reçue", "2-Accusé de réception", "3-En cours d'examen", "4-Entretien planifié",
                    "5-Entretien réalisé", "6-Vivier", "7-Proposition envoyée", "8-Recruté(e)",
                    "9-Refusée", "10-Désistement"],
    "StatutClos":  ["N", "N", "N", "N", "N", "N", "N", "O", "O", "O"],
    "Responsables": ["— à compléter —", "Associé 1", "Associé 2", "Responsable RH", "Secrétariat", "Non attribué"],
    "Evaluations": [1, 2, 3, 4, 5],
    "Decisions":   ["À statuer", "Favorable", "Vivier", "Défavorable"],
    "Motifs":      ["Profil non adapté au poste", "Expérience insuffisante", "Spécialité différente",
                    "Prétentions salariales", "Poste déjà pourvu", "Pas de disponibilité compatible",
                    "Sans suite du candidat", "Autre"],
    "Actions":     ["Accuser réception", "Étudier le dossier", "Demander pièces manquantes",
                    "Planifier un entretien", "Second entretien", "Prendre les références",
                    "Envoyer une proposition", "Envoyer un refus", "Relancer le candidat",
                    "Classer en vivier"],
    "OuiNon":      ["Oui", "Non"],
    "Capa":        ["Oui", "Non", "En cours", "Non concerné"],
    "TypesEntretien":   ["Au cabinet", "Visioconférence", "Téléphone"],
    "StatutsEntretien": ["Planifié", "Réalisé", "Annulé", "Reporté"],
    "AvisEntretien":    ["Très favorable", "Favorable", "Réservé", "Défavorable"],
}
ORDRE = ["Civilites", "Postes", "Sources", "StatutListe", "StatutClos", "Responsables", "Evaluations",
         "Decisions", "Motifs", "Actions", "OuiNon", "Capa", "TypesEntretien", "StatutsEntretien",
         "AvisEntretien"]
TITRES = {"Civilites": "Civilité", "Postes": "Postes", "Sources": "Sources", "StatutListe": "Statuts",
          "StatutClos": "Clôturé ?", "Responsables": "Responsables", "Evaluations": "Évaluation",
          "Decisions": "Décisions", "Motifs": "Motifs de refus", "Actions": "Prochaines actions",
          "OuiNon": "Oui / Non", "Capa": "CAPA / CRFPA", "TypesEntretien": "Type d'entretien",
          "StatutsEntretien": "Statut entretien", "AvisEntretien": "Avis entretien"}

LROW = 29
section(ws, LROW, "LISTES DÉROULANTES — ajoutez ou renommez librement les valeurs", 2, 12)
ws.cell(row=LROW + 1, column=2,
        value="Après ajout d'une valeur en fin de liste, cliquez sur « Actualiser les listes » "
              "(onglet Accueil) pour étendre le menu déroulant correspondant.").font = F(9, False, GREY, it=True)
HROW = LROW + 2
for i, key in enumerate(ORDRE):
    col = 2 + i
    letter = get_column_letter(col)
    h = ws.cell(row=HROW, column=col, value=TITRES[key])
    h.font = F(9, True, WHITE); h.fill = P(GREY); h.border = BOX
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    vals = LISTES[key]
    for j, v in enumerate(vals):
        c = ws.cell(row=HROW + 1 + j, column=col, value=v)
        c.font = F(9); c.border = BOX
        c.alignment = Alignment(horizontal="center" if key in ("Evaluations", "StatutClos", "OuiNon") else "left",
                                indent=0 if key in ("Evaluations", "StatutClos", "OuiNon") else 1)
    ws.column_dimensions[letter].width = max(11, min(26, len(TITRES[key]) + 6,
                                                     max(len(str(v)) for v in vals) + 4))
    wb.defined_names.add(DefinedName(key, attr_text="'Paramètres'!$%s$%d:$%s$%d"
                                     % (letter, HROW + 1, letter, HROW + len(vals))))
ws.row_dimensions[HROW].height = 28
ws.freeze_panes = "B4"
ws.sheet_properties.tabColor = GREY

# ================================================================ CANDIDATURES
wc = wb.create_sheet("Candidatures")
# (en-tête, largeur, type, format, liste de validation)
COLS = [
    ("ID",                    12, "auto", None,          None),
    ("Nom",                   18, "in",   None,          None),
    ("Prénom",                14, "in",   None,          None),
    ("Statut",                21, "in",   None,          "StatutListe"),
    ("Alerte",                13, "calc", None,          None),
    ("Civilité",               9, "in",   None,          "Civilites"),
    ("Date réception",        13, "in",   "dd/mm/yyyy",  None),
    ("E-mail",                28, "in",   None,          None),
    ("Téléphone",             15, "in",   "@",           None),
    ("Ville",                 14, "in",   None,          None),
    ("Poste visé",            25, "in",   None,          "Postes"),
    ("Source",                21, "in",   None,          "Sources"),
    ("Expérience (ans)",      10, "in",   "0.0",         None),
    ("Diplôme / École",       28, "in",   None,          None),
    ("CAPA / CRFPA",          12, "in",   None,          "Capa"),
    ("Disponibilité",         12, "in",   "dd/mm/yyyy",  None),
    ("Prétentions (€)",       13, "in",   '#,##0" €"',   None),
    ("Responsable",           16, "in",   None,          "Responsables"),
    ("Éval. /5",               9, "in",   '0" ★"',       "Evaluations"),
    ("Date entretien",        16, "in",   "dd/mm/yyyy hh:mm", None),
    ("Décision",              13, "in",   None,          "Decisions"),
    ("Motif (si refus)",      26, "in",   None,          "Motifs"),
    ("Date dernier contact",  14, "in",   "dd/mm/yyyy",  None),
    ("Prochaine action",      23, "in",   None,          "Actions"),
    ("Date prochaine action", 14, "in",   "dd/mm/yyyy",  None),
    ("Délai réponse (j)",     11, "calc", "0",           None),
    ("Ancienneté (j)",        11, "calc", "0",           None),
    ("Dossier candidat",      24, "in",   None,          None),
    ("Consentement vivier",   12, "in",   None,          "OuiNon"),
    ("Purge RGPD le",         13, "calc", "dd/mm/yyyy",  None),
    ("Commentaires",          45, "in",   None,          None),
]
NCOL = len(COLS)                       # 31 -> A..AE
LAST = get_column_letter(NCOL)
HDR, FIRST = 4, 5
LASTDATA = FIRST + NDATA - 1           # 34
DVEND = 1004                           # portée des listes / MFC au-delà du tableau

bandeau(wc, NCOL, "SUIVI DES CANDIDATURES",
        "Une ligne = une candidature.  Sélectionnez une ligne puis utilisez les boutons ci-dessous "
        "(ou l'onglet Accueil).  Colonnes grisées = calculées automatiquement, ne rien y saisir.")
wc.row_dimensions[2].height = 30       # emplacement des boutons créés par la macro Installer
for col in range(1, NCOL + 1):
    wc.cell(row=2, column=col).fill = P(BLUE_2)

GROUPES = [(1, 5, "SUIVI RAPIDE", NAVY), (6, 10, "COORDONNÉES", "2C5F8A"),
           (11, 17, "PROFIL & POSTE", "3C6E9E"), (18, 25, "TRAITEMENT DE LA CANDIDATURE", "2C5F8A"),
           (26, 27, "DÉLAIS (auto)", GREY), (28, 31, "DOSSIER & RGPD", GOLD)]
for c1, c2, lib, coul in GROUPES:
    wc.merge_cells(start_row=3, start_column=c1, end_row=3, end_column=c2)
    c = wc.cell(row=3, column=c1, value=lib)
    c.font = F(8, True, WHITE); c.fill = P(coul); c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")
wc.row_dimensions[3].height = 16

for i, (head, width, kind, fmt, dvname) in enumerate(COLS, start=1):
    letter = get_column_letter(i)
    wc.column_dimensions[letter].width = width
    h = wc.cell(row=HDR, column=i, value=head)
    h.font = F(9, True, WHITE)
    h.fill = P(GREY if kind == "calc" else NAVY)
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    h.border = BOX
    for r in range(FIRST, LASTDATA + 1):
        c = wc.cell(row=r, column=i)
        c.font = F(9.5)
        c.border = BOX
        if fmt:
            c.number_format = fmt
        if kind == "calc":
            c.fill = P(FILL_CALC)
        c.alignment = Alignment(
            horizontal="center" if (fmt or kind == "calc" or dvname in ("Civilites", "Capa", "OuiNon", "Evaluations")) else "left",
            vertical="center", wrap_text=(head == "Commentaires"), indent=0 if fmt else 1)
wc.row_dimensions[HDR].height = 30

# --- formules des colonnes calculées
CALC = {
    "E":  '=IF($A{r}="","",IF(IFERROR(INDEX(StatutClos,MATCH($D{r},StatutListe,0)),"N")="O","Clôturé",'
          'IF(AND($Y{r}<>"",$Y{r}<Aujourdhui),"Retard",IF($Y{r}=Aujourdhui,"Aujourd\'hui",'
          'IF(AND($W{r}="",$G{r}<>"",Aujourdhui-$G{r}>SeuilAccuse),"À traiter","OK")))))',
    "Z":  '=IF(OR($A{r}="",$G{r}="",$W{r}=""),"",$W{r}-$G{r})',
    "AA": '=IF(OR($A{r}="",$G{r}=""),"",Aujourdhui-$G{r})',
    "AD": '=IF($A{r}="","",IF($W{r}<>"",EDATE($W{r},12*DureeConservation),'
          'IF($G{r}="","",EDATE($G{r},12*DureeConservation))))',
}
for r in range(FIRST, LASTDATA + 1):
    for letter, tpl in CALC.items():
        wc["%s%d" % (letter, r)] = tpl.format(r=r)

# --- ligne d'exemple (à supprimer par le cabinet)
AUJ = dt.date.today()
EX = {
    "A": "EXEMPLE-01", "B": "MARTIN", "C": "Camille", "D": "4-Entretien planifié", "F": "Mme",
    "G": AUJ - dt.timedelta(days=12), "H": "camille.martin@exemple.fr", "I": "06 12 34 56 78",
    "J": "Lyon", "K": "Avocat collaborateur", "L": "Village de la Justice", "M": 3,
    "N": "M2 Droit du dommage corporel — Lyon III", "O": "Oui",
    "P": AUJ + dt.timedelta(days=45), "Q": 42000, "R": "Associé 1", "S": 4,
    "T": dt.datetime.combine(AUJ + dt.timedelta(days=5), dt.time(10, 0)), "U": "À statuer",
    "W": AUJ - dt.timedelta(days=10), "X": "Prendre les références",
    "Y": AUJ + dt.timedelta(days=3), "AB": r"C:\Cabinet\Recrutement\Candidats\EXEMPLE-01_MARTIN_Camille",
    "AC": "Oui",
    "AE": "LIGNE D'EXEMPLE — supprimez-la avant la première utilisation "
          "(bouton « Supprimer les exemples » de l'onglet Accueil).",
}
for letter, val in EX.items():
    wc["%s%d" % (letter, FIRST)] = val

tab = Table(displayName="tblCandidatures", ref="A%d:%s%d" % (HDR, LAST, LASTDATA))
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True,
                                    showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
wc.add_table(tab)

# --- listes déroulantes
def add_dv(sheet, listname, ranges, strict=False, msg=""):
    dv = DataValidation(type="list", formula1=listname, allow_blank=True,
                        showErrorMessage=True, errorStyle="stop" if strict else "warning")
    dv.errorTitle = "Valeur hors liste"
    dv.error = msg or ("Choisissez une valeur dans la liste déroulante.\n"
                       "Pour ajouter une valeur, complétez la liste dans l'onglet Paramètres.")
    dv.promptTitle = ""
    sheet.add_data_validation(dv)
    for rg in ranges:
        dv.add(rg)
    return dv

for i, (head, width, kind, fmt, dvname) in enumerate(COLS, start=1):
    if not dvname:
        continue
    letter = get_column_letter(i)
    add_dv(wc, dvname, ["%s%d:%s%d" % (letter, FIRST, letter, DVEND)],
           strict=(dvname == "StatutListe"),
           msg=("Le statut pilote les indicateurs et les automatisations : "
                "choisissez-en un dans la liste." if dvname == "StatutListe" else ""))

dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
                     allow_blank=True, showErrorMessage=True, errorStyle="warning")
dvd.errorTitle = "Date invalide"
dvd.error = "Saisissez une date au format jj/mm/aaaa."
wc.add_data_validation(dvd)
for letter in ("G", "P", "W", "Y"):
    dvd.add("%s%d:%s%d" % (letter, FIRST, letter, DVEND))

# --- mises en forme conditionnelles
def mfc(sheet, rng, formule, fill=None, color=None, bold=False, italic=False):
    sheet.conditional_formatting.add(
        rng, FormulaRule(formula=[formule], stopIfTrue=False,
                         fill=P(fill) if fill else None,
                         font=Font(name="Arial", size=9.5, bold=bold, italic=italic,
                                   color=color or INK)))

STATUT_COUL = {1: ("E8EEF5", NAVY, False), 2: ("D6E4F0", NAVY, False), 3: ("FCF3CF", "7D6608", False),
               4: ("FDEBD0", "9C640C", True), 5: ("FAE5D3", "9C640C", False), 6: ("E8DAEF", "633974", False),
               7: ("D4EFDF", "196F3D", True), 8: ("ABEBC6", "145A32", True), 9: ("EDEDED", GREY, False),
               10: ("EDEDED", GREY, False)}
rng_statut = "D%d:D%d" % (FIRST, DVEND)
for n, (bg, fg, bold) in STATUT_COUL.items():
    mfc(wc, rng_statut, '=$D%d=INDEX(StatutListe,%d)' % (FIRST, n), bg, fg, bold)

rng_alerte = "E%d:E%d" % (FIRST, DVEND)
for val, bg, fg, bold in (("Retard", "F5B7B1", "922B21", True), ("Aujourd'hui", "FAD7A0", "9C640C", True),
                          ("À traiter", "FCF3CF", "7D6608", False), ("OK", "E9F7EF", OK_G, False),
                          ("Clôturé", "F2F3F4", GREY, False)):
    mfc(wc, rng_alerte, '=$E%d="%s"' % (FIRST, val), bg, fg, bold)

# lignes clôturées en gris sur toute la largeur
mfc(wc, "A%d:%s%d" % (FIRST, LAST, DVEND), '=$E%d="Clôturé"' % FIRST, None, GREY, italic=True)
# purge RGPD échue
mfc(wc, "AD%d:AD%d" % (FIRST, DVEND), '=AND($AD%d<>"",$AD%d<=Aujourdhui)' % (FIRST, FIRST),
    "F5B7B1", "922B21", True)
# ancienneté
wc.conditional_formatting.add("AA%d:AA%d" % (FIRST, DVEND),
    ColorScaleRule(start_type="num", start_value=0, start_color="E9F7EF",
                   mid_type="num", mid_value=30, mid_color="FCF3CF",
                   end_type="num", end_value=90, end_color="F5B7B1"))
# évaluation
wc.conditional_formatting.add("S%d:S%d" % (FIRST, DVEND),
    ColorScaleRule(start_type="num", start_value=1, start_color="F5B7B1",
                   mid_type="num", mid_value=3, mid_color="FCF3CF",
                   end_type="num", end_value=5, end_color="ABEBC6"))

wc.freeze_panes = "F%d" % FIRST
wc.sheet_view.showGridLines = False
wc.sheet_properties.tabColor = NAVY
wc.page_setup.orientation = "landscape"
wc.page_setup.fitToWidth = 1
wc.page_setup.fitToHeight = 0
wc.sheet_properties.pageSetUpPr.fitToPage = True
wc.print_title_rows = "%d:%d" % (HDR, HDR)

# ================================================================ ENTRETIENS
we = wb.create_sheet("Entretiens")
ECOLS = [("N° entretien", 13, "auto", None, None), ("ID candidat", 13, "in", None, None),
         ("Candidat", 22, "calc", None, None), ("Poste visé", 24, "calc", None, None),
         ("Date", 12, "in", "dd/mm/yyyy", None), ("Heure", 9, "in", "hh:mm", None),
         ("Durée (min)", 10, "in", "0", None), ("Type", 16, "in", None, "TypesEntretien"),
         ("Lieu / lien visio", 30, "in", None, None), ("Intervenants du cabinet", 26, "in", None, None),
         ("Statut", 13, "in", None, "StatutsEntretien"), ("Note /20", 9, "in", "0.0", None),
         ("Points forts", 34, "in", None, None), ("Points de vigilance", 34, "in", None, None),
         ("Avis", 16, "in", None, "AvisEntretien"), ("Compte rendu rédigé", 12, "in", None, "OuiNon")]
NE = len(ECOLS); LASTE = get_column_letter(NE); NEDATA = 20; ELAST = FIRST + NEDATA - 1
bandeau(we, NE, "ENTRETIENS",
        "Alimenté automatiquement par le bouton « Planifier un entretien ». "
        "Le nom et le poste sont récupérés depuis l'onglet Candidatures à partir de l'ID.")
we.row_dimensions[2].height = 6
for i, (head, width, kind, fmt, dvname) in enumerate(ECOLS, start=1):
    letter = get_column_letter(i)
    we.column_dimensions[letter].width = width
    h = we.cell(row=3, column=i, value=head)
    h.font = F(9, True, WHITE); h.fill = P(GREY if kind == "calc" else NAVY)
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); h.border = BOX
    for r in range(4, ELAST + 1):
        c = we.cell(row=r, column=i)
        c.font = F(9.5); c.border = BOX
        if fmt:
            c.number_format = fmt
        if kind == "calc":
            c.fill = P(FILL_CALC)
        c.alignment = Alignment(horizontal="center" if fmt else "left", vertical="center",
                                wrap_text=head in ("Points forts", "Points de vigilance"), indent=0 if fmt else 1)
    if dvname:
        add_dv(we, dvname, ["%s4:%s500" % (letter, letter)])
we.row_dimensions[3].height = 30
for r in range(4, ELAST + 1):
    we["C%d" % r] = ('=IF($B{r}="","",IFERROR(INDEX(Candidatures!$B$5:$B${m},MATCH($B{r},Candidatures!$A$5:$A${m},0))'
                     '&" "&INDEX(Candidatures!$C$5:$C${m},MATCH($B{r},Candidatures!$A$5:$A${m},0)),"ID inconnu"))').format(r=r, m=MAXROW)
    we["D%d" % r] = ('=IF($B{r}="","",IFERROR(INDEX(Candidatures!$K$5:$K${m},MATCH($B{r},Candidatures!$A$5:$A${m},0)),""))').format(r=r, m=MAXROW)
we["A4"], we["B4"], we["E4"] = "EXEMPLE-E01", "EXEMPLE-01", AUJ + dt.timedelta(days=5)
we["F4"], we["G4"], we["H4"] = dt.time(10, 0), 45, "Au cabinet"
we["I4"], we["J4"], we["K4"] = "Salle de réunion — 2e étage", "Associé 1, Responsable RH", "Planifié"
we["M4"] = "Exemple : maîtrise de la nomenclature Dintilhac."
tabe = Table(displayName="tblEntretiens", ref="A3:%s%d" % (LASTE, ELAST))
tabe.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True, showColumnStripes=False,
                                     showFirstColumn=False, showLastColumn=False)
we.add_table(tabe)
we.freeze_panes = "E4"
we.sheet_view.showGridLines = False
we.sheet_properties.tabColor = NAVY
we.page_setup.orientation = "landscape"

# ================================================================ JOURNAL
wj = wb.create_sheet("Journal")
JCOLS = [("Horodatage", 18, "dd/mm/yyyy hh:mm"), ("Utilisateur", 18, None), ("ID candidat", 14, None),
         ("Candidat", 24, None), ("Action", 26, None), ("Détail", 60, None), ("Résultat", 14, None)]
bandeau(wj, len(JCOLS), "JOURNAL DES ACTIONS",
        "Rempli automatiquement par les macros (e-mails envoyés, documents générés, imports). "
        "Sert de preuve de traitement et de traçabilité RGPD. Ne pas saisir manuellement.")
wj.row_dimensions[2].height = 6
for i, (head, width, fmt) in enumerate(JCOLS, start=1):
    letter = get_column_letter(i)
    wj.column_dimensions[letter].width = width
    h = wj.cell(row=3, column=i, value=head)
    h.font = F(9, True, WHITE); h.fill = P(NAVY); h.border = BOX
    h.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(4, 24):
        c = wj.cell(row=r, column=i); c.font = F(9); c.border = BOX
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="center" if fmt else "left", vertical="center", indent=0 if fmt else 1)
wj.row_dimensions[3].height = 22
wj["A4"], wj["B4"], wj["C4"] = dt.datetime.now().replace(second=0, microsecond=0), "Exemple", "EXEMPLE-01"
wj["D4"], wj["E4"] = "MARTIN Camille", "Accusé de réception"
wj["F4"], wj["G4"] = "E-mail envoyé à camille.martin@exemple.fr", "OK"
tabj = Table(displayName="tblJournal", ref="A3:%s23" % get_column_letter(len(JCOLS)))
tabj.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True, showColumnStripes=False,
                                     showFirstColumn=False, showLastColumn=False)
wj.add_table(tabj)
wj.freeze_panes = "A4"
wj.sheet_view.showGridLines = False
wj.sheet_properties.tabColor = GREY

# ================================================================ MODÈLES E-MAILS
wm = wb.create_sheet("Modèles e-mails")
SIG = "\n\n{{SIGNATURE_BLOC}}"
RGPD = ("\n\nConformément au règlement (UE) 2016/679, les données de votre candidature sont conservées "
        "{{DUREE_CONSERVATION}} ans à compter de notre dernier échange, aux seules fins du recrutement. "
        "Vous pouvez exercer vos droits d'accès, de rectification et d'effacement à l'adresse {{EMAIL_RH}}.")
MODELES = [
 ("ACCUSE_RECEPTION", "Accusé de réception",
  "Votre candidature au poste de {{POSTE}} — {{CABINET}}",
  "{{CIVILITE_LONGUE}},\n\n"
  "Nous accusons réception de votre candidature au poste de {{POSTE}}, qui nous est parvenue le "
  "{{DATE_RECEPTION}}, et vous remercions de l'intérêt que vous portez à notre cabinet.\n\n"
  "Votre dossier est en cours d'examen par {{RESPONSABLE}}. Nous reviendrons vers vous dans un délai "
  "de trois semaines, que la suite donnée soit favorable ou non.\n\n"
  "Nous vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + RGPD + SIG, ""),
 ("DEMANDE_PIECES", "Demande de pièces complémentaires",
  "Votre candidature — pièces complémentaires",
  "{{CIVILITE_LONGUE}},\n\n"
  "Nous avons bien reçu votre candidature au poste de {{POSTE}}. Afin de poursuivre son examen, "
  "nous vous remercions de nous transmettre les éléments suivants :\n\n"
  "  •  votre curriculum vitae actualisé ;\n"
  "  •  une lettre de motivation ;\n"
  "  •  le cas échéant, une copie de votre CAPA ou de votre attestation de réussite ;\n"
  "  •  vos dates de disponibilité.\n\n"
  "Ces pièces peuvent être adressées en réponse au présent message.\n\n"
  "Nous vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + SIG, ""),
 ("CONVOCATION", "Convocation à un entretien",
  "Entretien — poste de {{POSTE}} — {{CABINET}}",
  "{{CIVILITE_LONGUE}},\n\n"
  "Votre candidature au poste de {{POSTE}} a retenu notre attention.\n\n"
  "Nous vous proposons de nous rencontrer le {{DATE_ENTRETIEN}} à {{HEURE_ENTRETIEN}}.\n"
  "Modalité : {{TYPE_ENTRETIEN}}\n"
  "Lieu / lien : {{LIEU_ENTRETIEN}}\n"
  "Durée prévue : environ {{DUREE_ENTRETIEN}} minutes\n"
  "Vous serez reçu(e) par {{RESPONSABLE}}.\n\n"
  "Nous vous remercions de bien vouloir confirmer votre présence en réponse à ce message. "
  "Si ce créneau ne vous convenait pas, n'hésitez pas à nous proposer d'autres disponibilités.\n\n"
  "Nous vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + SIG, ""),
 ("RELANCE_CANDIDAT", "Relance du candidat",
  "Votre candidature au poste de {{POSTE}} — relance",
  "{{CIVILITE_LONGUE}},\n\n"
  "Sauf erreur de notre part, notre message du {{DATE_DERNIER_CONTACT}} relatif à votre candidature "
  "au poste de {{POSTE}} est resté sans réponse.\n\n"
  "Nous restons à votre disposition et vous remercions de nous indiquer si vous souhaitez maintenir "
  "votre candidature.\n\n"
  "Nous vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + SIG, ""),
 ("VIVIER", "Mise en vivier",
  "Votre candidature — {{CABINET}}",
  "{{CIVILITE_LONGUE}},\n\n"
  "Nous avons étudié avec attention votre candidature au poste de {{POSTE}}.\n\n"
  "Si nous ne sommes pas en mesure d'y donner une suite favorable dans l'immédiat, votre profil a "
  "retenu notre intérêt. Sous réserve de votre accord, nous conserverons votre dossier afin de vous "
  "recontacter dès qu'un poste correspondant se libérera.\n\n"
  "Il vous suffit de nous répondre par la négative pour que votre dossier soit supprimé sans délai.\n\n"
  "Nous vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + RGPD + SIG, ""),
 ("REFUS", "Réponse négative",
  "Votre candidature au poste de {{POSTE}} — {{CABINET}}",
  "{{CIVILITE_LONGUE}},\n\n"
  "Nous avons étudié avec attention votre candidature au poste de {{POSTE}} et vous remercions de la "
  "confiance que vous avez témoignée à notre cabinet.\n\n"
  "Nous sommes au regret de ne pouvoir y donner une suite favorable. Cette décision ne remet nullement "
  "en cause la qualité de votre parcours ; elle tient à l'adéquation entre votre profil et les besoins "
  "actuels du cabinet.\n\n"
  "Nous vous souhaitons une pleine réussite dans la suite de votre parcours professionnel et vous prions "
  "d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos salutations distinguées." + RGPD + SIG, ""),
 ("PROPOSITION", "Proposition / suite favorable",
  "Votre candidature au poste de {{POSTE}} — suite favorable",
  "{{CIVILITE_LONGUE}},\n\n"
  "À l'issue de nos échanges, nous avons le plaisir de vous confirmer notre souhait de vous accueillir "
  "au sein du cabinet au poste de {{POSTE}}.\n\n"
  "Vous trouverez ci-joint le document reprenant les conditions envisagées. Nous vous remercions de "
  "nous faire part de vos observations et de votre date de disponibilité.\n\n"
  "Nous restons à votre entière disposition et vous prions d'agréer, {{CIVILITE_LONGUE}}, l'expression de nos "
  "salutations distinguées." + SIG, ""),
]
bandeau(wm, 5, "MODÈLES D'E-MAILS",
        "Textes envoyés par les boutons d'automatisation. Modifiez-les librement : "
        "ne touchez pas à la colonne « Clé », utilisée par les macros.")
wm.row_dimensions[2].height = 6
MHDR = ["Clé", "Libellé du bouton", "Objet de l'e-mail", "Corps du message", "Pièce jointe (chemin complet)"]
for i, h in enumerate(MHDR, start=1):
    c = wm.cell(row=3, column=i, value=h)
    c.font = F(9, True, WHITE); c.fill = P(NAVY); c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")
for w, letter in zip((20, 26, 46, 100, 30), "ABCDE"):
    wm.column_dimensions[letter].width = w
for r, (cle, lib, obj, corps, pj) in enumerate(MODELES, start=4):
    for i, v in enumerate((cle, lib, obj, corps, pj), start=1):
        c = wm.cell(row=r, column=i, value=v)
        c.font = F(9, True if i == 1 else False, NAVY if i == 1 else INK)
        c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=(i == 4), indent=1)
    wm.row_dimensions[r].height = 96
wm.row_dimensions[3].height = 22
tabm = Table(displayName="tblModeles", ref="A3:E%d" % (3 + len(MODELES)))
tabm.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True, showColumnStripes=False,
                                     showFirstColumn=False, showLastColumn=False)
wm.add_table(tabm)
AIDE = 3 + len(MODELES) + 2
wm.cell(row=AIDE, column=1, value="BALISES DISPONIBLES — remplacées automatiquement à l'envoi").font = F(10, True, NAVY)
BALISES = [
 ("{{CIVILITE}}", "M. / Mme — pour le bloc adresse"),
 ("{{CIVILITE_LONGUE}}", "Madame / Monsieur — pour l'appel et la formule de politesse"), ("{{NOM}}", "Nom du candidat"), ("{{PRENOM}}", "Prénom"),
 ("{{POSTE}}", "Poste visé"), ("{{ID}}", "Identifiant de la candidature"),
 ("{{DATE_RECEPTION}}", "Date de réception de la candidature"),
 ("{{DATE_DERNIER_CONTACT}}", "Date du dernier contact"),
 ("{{DATE_ENTRETIEN}}", "Date de l'entretien (jj/mm/aaaa)"), ("{{HEURE_ENTRETIEN}}", "Heure de l'entretien"),
 ("{{DUREE_ENTRETIEN}}", "Durée en minutes"), ("{{TYPE_ENTRETIEN}}", "Au cabinet / visio / téléphone"),
 ("{{LIEU_ENTRETIEN}}", "Adresse ou lien de connexion"), ("{{RESPONSABLE}}", "Responsable du dossier"),
 ("{{CABINET}}", "Nom du cabinet (Paramètres)"), ("{{ADRESSE}}", "Adresse du cabinet"),
 ("{{TELEPHONE}}", "Téléphone du cabinet"), ("{{EMAIL_RH}}", "E-mail de recrutement"),
 ("{{SITE}}", "Site internet"), ("{{SIGNATAIRE}}", "Nom du signataire"), ("{{FONCTION}}", "Fonction"),
 ("{{SIGNATURE_BLOC}}", "Bloc de signature complet, construit depuis Paramètres "
                       "(laissez « Signature e-mail » vide pour utiliser celle d'Outlook)"),
 ("{{DUREE_CONSERVATION}}", "Durée de conservation RGPD, en années"),
]
for k, (bal, desc) in enumerate(BALISES):
    r = AIDE + 1 + k
    a = wm.cell(row=r, column=1, value=bal); a.font = F(9, True, GOLD)
    b = wm.cell(row=r, column=2, value=desc); b.font = F(9, False, GREY)
wm.sheet_view.showGridLines = False
wm.sheet_properties.tabColor = GOLD

# ================================================================ TABLEAU DE BORD
wd = wb.create_sheet("Tableau de bord")
bandeau(wd, 12, "TABLEAU DE BORD DU RECRUTEMENT",
        "Tous les chiffres se recalculent automatiquement à chaque saisie. Aucune manipulation requise.")
wd.column_dimensions["A"].width = 2
for letter in "BCDEFGHIJK":
    wd.column_dimensions[letter].width = 14
wd.column_dimensions["L"].width = 2
C = "Candidatures!"
TOT = "COUNTA(%s$A$5:$A$%d)" % (C, MAXROW)
def card(row, col, libelle, formule, fmt="0", couleur=NAVY):
    wd.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    l = wd.cell(row=row, column=col, value=libelle)
    l.font = F(8.5, True, WHITE); l.fill = P(couleur); l.border = BOX
    l.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wd.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    v = wd.cell(row=row + 1, column=col, value=formule)
    v.font = F(20, True, couleur); v.fill = P(BLUE_2); v.border = BOX
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.number_format = fmt
    wd.row_dimensions[row].height = 26
    wd.row_dimensions[row + 1].height = 34

section(wd, 4, "INDICATEURS CLÉS", 2, 9)
STATUT_N = lambda n: "INDEX(StatutListe,%d)" % n
KPI1 = [
  ("Candidatures reçues", "=%s" % TOT, "0", NAVY),
  ("Dossiers en cours", "=%s-COUNTIF(%s$D$5:$D$%d,%s)-COUNTIF(%s$D$5:$D$%d,%s)-COUNTIF(%s$D$5:$D$%d,%s)"
   % (TOT, C, MAXROW, STATUT_N(8), C, MAXROW, STATUT_N(9), C, MAXROW, STATUT_N(10)), "0", "2C5F8A"),
  ("Entretiens à venir", '=COUNTIFS(%s$T$5:$T$%d,">="&Aujourdhui)' % (C, MAXROW), "0", "9C640C"),
  ("Recrutements", "=COUNTIF(%s$D$5:$D$%d,%s)" % (C, MAXROW, STATUT_N(8)), "0", OK_G),
  ("Réponses négatives", "=COUNTIF(%s$D$5:$D$%d,%s)" % (C, MAXROW, STATUT_N(9)), "0", GREY),
]
for k, (lib, f, fmt, coul) in enumerate(KPI1):
    card(5, 2 + 2 * k, lib, f, fmt, coul)
KPI2 = [
  ("À traiter", '=COUNTIF(%s$E$5:$E$%d,"À traiter")' % (C, MAXROW), "0", "C87F0A"),
  ("Actions en retard", '=COUNTIF(%s$E$5:$E$%d,"Retard")' % (C, MAXROW), "0", BAD_R),
  ("Délai moyen de réponse", "=IFERROR(AVERAGE(%s$Z$5:$Z$%d),0)" % (C, MAXROW), '0.0" j"', NAVY),
  ("Taux d'entretien", "=IFERROR((COUNTIF(%s$D$5:$D$%d,%s)+COUNTIF(%s$D$5:$D$%d,%s)+COUNTIF(%s$D$5:$D$%d,%s)"
   "+COUNTIF(%s$D$5:$D$%d,%s))/%s,0)" % (C, MAXROW, STATUT_N(4), C, MAXROW, STATUT_N(5), C, MAXROW,
                                          STATUT_N(7), C, MAXROW, STATUT_N(8), TOT), "0.0%", "2C5F8A"),
  ("À purger (RGPD)", '=COUNTIFS(%s$AD$5:$AD$%d,">0",%s$AD$5:$AD$%d,"<="&Aujourdhui)'
   % (C, MAXROW, C, MAXROW), "0", GOLD),
]
for k, (lib, f, fmt, coul) in enumerate(KPI2):
    card(8, 2 + 2 * k, lib, f, fmt, coul)

def bloc(row, col, titre, nom_liste, n, col_source):
    wd.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 3)
    t = wd.cell(row=row, column=col, value=titre)
    t.font = F(10, True, WHITE); t.fill = P(NAVY); t.border = BOX
    t.alignment = Alignment(vertical="center", indent=1)
    wd.row_dimensions[row].height = 20
    for j, h in enumerate(("Libellé", "Nb", "Part", "")):
        c = wd.cell(row=row + 1, column=col + j, value=h)
        c.font = F(8.5, True, GREY); c.border = BOX
        c.alignment = Alignment(horizontal="center" if j else "left", vertical="center", indent=0 if j else 1)
    for i in range(1, n + 1):
        r = row + 1 + i
        lb = get_column_letter(col)
        a = wd.cell(row=r, column=col, value="=INDEX(%s,%d)" % (nom_liste, i))
        a.font = F(9); a.border = BOX; a.alignment = Alignment(indent=1)
        b = wd.cell(row=r, column=col + 1, value="=COUNTIF(%s$%s$5:$%s$%d,$%s%d)"
                    % (C, col_source, col_source, MAXROW, lb, r))
        b.font = F(9, True); b.border = BOX; b.alignment = Alignment(horizontal="center")
        d = wd.cell(row=r, column=col + 2, value="=IFERROR(%s/%s,0)" % (b.coordinate, TOT))
        d.font = F(9, False, GREY); d.border = BOX; d.number_format = "0.0%"
        d.alignment = Alignment(horizontal="center")
        e = wd.cell(row=r, column=col + 3, value="=%s" % b.coordinate)
        e.font = F(9, False, WHITE); e.border = BOX
    rng = "%s%d:%s%d" % (get_column_letter(col + 3), row + 2, get_column_letter(col + 3), row + 1 + n)
    wd.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="max",
                                                   color=NAVY, showValue=False))
    wd.column_dimensions[get_column_letter(col + 3)].width = 16
    return row + 2 + n

r1 = bloc(12, 2, "RÉPARTITION PAR STATUT", "StatutListe", 10, "D")
r2 = bloc(12, 7, "RÉPARTITION PAR POSTE", "Postes", 10, "K")
r3 = bloc(max(r1, r2) + 1, 2, "RÉPARTITION PAR SOURCE", "Sources", 10, "L")
r4 = bloc(max(r1, r2) + 1, 7, "RÉPARTITION PAR RESPONSABLE", "Responsables", 6, "R")

RM = max(r3, r4) + 1
wd.merge_cells(start_row=RM, start_column=2, end_row=RM, end_column=10)
t = wd.cell(row=RM, column=2, value="FLUX DES CANDIDATURES — 12 DERNIERS MOIS")
t.font = F(10, True, WHITE); t.fill = P(NAVY); t.border = BOX
t.alignment = Alignment(vertical="center", indent=1)
wd.row_dimensions[RM].height = 20
for i in range(12):
    r = RM + 1 + i
    a = wd.cell(row=r, column=2, value="=EDATE(DATE(YEAR(Aujourdhui),MONTH(Aujourdhui),1),%d)" % (i - 11))
    a.number_format = "mmmm yyyy"; a.font = F(9); a.border = BOX; a.alignment = Alignment(indent=1)
    b = wd.cell(row=r, column=3, value='=COUNTIFS(%s$G$5:$G$%d,">="&$B%d,%s$G$5:$G$%d,"<"&EDATE($B%d,1))'
                % (C, MAXROW, r, C, MAXROW, r))
    b.font = F(9, True); b.border = BOX; b.alignment = Alignment(horizontal="center")
    e = wd.cell(row=r, column=4, value="=$C%d" % r); e.font = F(9, False, WHITE); e.border = BOX
    wd.merge_cells(start_row=r, start_column=4, end_row=r, end_column=10)
wd.conditional_formatting.add("D%d:D%d" % (RM + 1, RM + 12),
                              DataBarRule(start_type="num", start_value=0, end_type="max",
                                          color="2C5F8A", showValue=False))
NOTE = RM + 14
wd.cell(row=NOTE, column=2,
        value="Périmètre de calcul : lignes 5 à %d de l'onglet Candidatures. "
              "Les libellés proviennent des listes de l'onglet Paramètres : renommer une valeur "
              "met automatiquement à jour ce tableau." % MAXROW).font = F(8.5, False, GREY, it=True)
wd.cell(row=NOTE + 1, column=2,
        value="Durée de conservation des candidatures non retenues : paramétrable dans l'onglet "
              "Paramètres — recommandation CNIL : 2 ans à compter du dernier contact.").font = F(8.5, False, GREY, it=True)
wd.sheet_view.showGridLines = False
wd.sheet_properties.tabColor = GOLD

# ================================================================ ACCUEIL
wa = wb.create_sheet("Accueil")
wa.column_dimensions["A"].width = 2
for letter, w in (("B", 34), ("C", 4), ("D", 34), ("E", 4), ("F", 3),
                  ("G", 30), ("H", 46), ("I", 14), ("J", 14), ("K", 14)):
    wa.column_dimensions[letter].width = w
bandeau(wa, 11, "GESTION DES CANDIDATURES",
        "Cabinet Victimes & Préjudices — tableau de bord, suivi des candidats, automatisations Outlook et Word")

D = "'Tableau de bord'!"
KPI_ACC = [("Candidatures", D + "B6", NAVY), ("En cours", D + "D6", "2C5F8A"),
           ("Entretiens à venir", D + "F6", "9C640C"), ("À traiter", D + "B9", "C87F0A"),
           ("Actions en retard", D + "D9", BAD_R)]
for k, (lib, ref, coul) in enumerate(KPI_ACC):
    col = 2 + 2 * k
    wa.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    c = wa.cell(row=4, column=col, value=lib)
    c.font = F(8.5, True, WHITE); c.fill = P(coul); c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wa.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
    v = wa.cell(row=5, column=col, value="=%s" % ref)
    v.font = F(22, True, coul); v.fill = P(BLUE_2); v.border = BOX
    v.alignment = Alignment(horizontal="center", vertical="center")
wa.row_dimensions[4].height = 24
wa.row_dimensions[5].height = 40

wa.merge_cells("B7:K7")
c = wa["B7"]
c.value = ("Ces compteurs se mettent à jour automatiquement.  Détail complet : onglet « Tableau de bord ».")
c.font = F(9, False, GREY, it=True)

section(wa, 9, "PILOTAGE  —  sélectionnez une ligne dans l'onglet Candidatures, puis cliquez", 2, 3)
section(wa, 9, "NAVIGATION ET AIDE", 7, 1)
wa.merge_cells("B10:E10")
z = wa["B10"]
z.value = ("Les boutons s'affichent ici après l'installation du module d'automatisation "
           "(onglet Notice, 5 minutes). Sans macros, le classeur reste pleinement utilisable à la main.")
z.font = F(9, False, GREY, it=True)
z.alignment = Alignment(wrap_text=True, vertical="top")
wa.row_dimensions[10].height = 30
for r in range(11, 30):
    for col in range(2, 6):
        wa.cell(row=r, column=col).fill = P(BLUE_2)
    wa.row_dimensions[r].height = 16
wb.defined_names.add(DefinedName("ZoneBoutons", attr_text="'Accueil'!$B$11"))

NAV = [("Candidatures", "Candidatures", "Le tableau de suivi : une ligne par candidature."),
       ("Entretiens", "Entretiens", "Planning et comptes rendus d'entretien."),
       ("Tableau de bord", "Tableau de bord", "Statistiques, délais, flux mensuel."),
       ("Journal", "Journal", "Historique automatique des e-mails et documents."),
       ("Modèles e-mails", "Modèles e-mails", "Textes des messages types, librement modifiables."),
       ("Paramètres", "Paramètres", "Coordonnées, dossiers, délais, listes déroulantes."),
       ("Notice", "Notice", "Mode d'emploi et installation des automatisations.")]
for k, (lib, sheet, desc) in enumerate(NAV):
    r = 11 + k
    lien(wa, "G%d" % r, sheet, "▸  " + lib)
    wa.cell(row=r, column=8, value=desc).font = F(9, False, GREY)
    wa.row_dimensions[r].height = 16

wa.cell(row=19, column=7, value="AVANT LA PREMIÈRE UTILISATION").font = F(10, True, NAVY)
ETAPES = ["1.  Onglet Paramètres : renseignez les coordonnées du cabinet et les chemins des dossiers.",
          "2.  Onglet Paramètres : adaptez les listes (postes, responsables, sources) à votre organisation.",
          "3.  Onglet Modèles e-mails : relisez et ajustez les textes types.",
          "4.  Onglet Notice : installez le module d'automatisation (Outlook / Word).",
          "5.  Onglet Candidatures : supprimez la ligne EXEMPLE-01, puis saisissez vos candidatures."]
for k, t in enumerate(ETAPES):
    r = 20 + k
    wa.merge_cells(start_row=r, start_column=7, end_row=r, end_column=11)
    c = wa.cell(row=r, column=7, value=t)
    c.font = F(9); c.alignment = Alignment(indent=1, vertical="center")
    wa.row_dimensions[r].height = 16

wa.cell(row=26, column=7, value="BON À SAVOIR").font = F(10, True, NAVY)
ASTUCES = ["Les colonnes sur fond gris sont calculées : ne rien y saisir.",
           "La colonne « Alerte » signale les dossiers en retard ou à traiter.",
           "Pour ajouter une candidature : cliquez sur la première ligne vide sous le tableau.",
           "Filtrez et triez avec les flèches des en-têtes du tableau."]
for k, t in enumerate(ASTUCES):
    r = 27 + k
    wa.merge_cells(start_row=r, start_column=7, end_row=r, end_column=11)
    c = wa.cell(row=r, column=7, value="•  " + t)
    c.font = F(9, False, GREY); c.alignment = Alignment(indent=1, vertical="center")
    wa.row_dimensions[r].height = 15
wa.sheet_view.showGridLines = False
wa.sheet_properties.tabColor = GOLD

# ================================================================ NOTICE
wn = wb.create_sheet("Notice")
wn.column_dimensions["A"].width = 2
wn.column_dimensions["B"].width = 34
wn.column_dimensions["C"].width = 108
bandeau(wn, 3, "MODE D'EMPLOI", "Prise en main, installation des automatisations Outlook et Word, dépannage")
LIGNES = []
def S(t):  LIGNES.append(("S", t, ""))
def L(a, b=""): LIGNES.append(("L", a, b))
def N(t):  LIGNES.append(("N", "", t))
def K(t):  LIGNES.append(("K", "", t))

S("1.  À QUOI SERT CHAQUE ONGLET")
L("Accueil", "Compteurs du jour, boutons d'automatisation, liens vers les autres onglets.")
L("Candidatures", "Le cœur de l'outil : une ligne par candidature reçue, de la réception à la décision.")
L("Entretiens", "Un entretien par ligne : date, intervenants, note, avis. Rempli par le bouton de planification.")
L("Tableau de bord", "Statistiques automatiques : statuts, postes, sources, responsables, délais, flux mensuel.")
L("Journal", "Trace horodatée de chaque e-mail envoyé et de chaque document généré.")
L("Modèles e-mails", "Les textes envoyés par les boutons. Modifiables librement, sauf la colonne « Clé ».")
L("Paramètres", "Coordonnées du cabinet, dossiers de travail, délais, et toutes les listes déroulantes.")
L("Notice", "Le présent mode d'emploi.")

S("2.  PRISE EN MAIN EN CINQ MINUTES")
L("Étape 1", "Onglet Paramètres : complétez les cellules sur fond crème (adresse, téléphone, e-mail de recrutement, signataire).")
L("Étape 2", "Toujours dans Paramètres : indiquez les trois dossiers de travail (candidats, modèles Word, dossier Outlook).")
L("Étape 3", "Adaptez les listes déroulantes à votre organisation : postes ouverts, responsables, sources de candidature.")
L("Étape 4", "Onglet Candidatures : supprimez la ligne EXEMPLE-01, puis saisissez votre première candidature.")
L("Étape 5", "Installez le module d'automatisation (section 3) pour activer les boutons Outlook et Word.")
N("Le classeur fonctionne parfaitement sans macros : vous perdez seulement l'envoi automatique des e-mails et la génération des documents Word.")

S("3.  ACTIVER LES AUTOMATISATIONS OUTLOOK ET WORD")
N("À faire une seule fois, sur chaque poste qui utilisera les boutons. Durée : environ 5 minutes.")
L("1.  Enregistrer au bon format",
  "Fichier ▸ Enregistrer sous ▸ type « Classeur Excel prenant en charge les macros (*.xlsm) ». "
  "Indispensable : un fichier .xlsx ne peut pas contenir de macros.")
L("2.  Débloquer le fichier",
  "Fermez Excel. Dans l'Explorateur, clic droit sur le fichier ▸ Propriétés ▸ cochez « Débloquer » ▸ OK. "
  "Windows bloque par défaut les macros des fichiers reçus par e-mail ou téléchargés.")
L("3.  Ouvrir l'éditeur de macros", "Rouvrez le classeur, puis appuyez sur Alt + F11.")
L("4.  Importer le module",
  "Dans l'éditeur : menu Fichier ▸ Importer un fichier… (raccourci Ctrl + M) ▸ sélectionnez "
  "VP_Candidatures.bas ▸ Ouvrir. Un élément « VP_Candidatures » apparaît sous Modules, "
  "à gauche, dans l'arborescence du classeur.")
L("4 bis.  Si l'import échoue",
  "Utilisez la variante copier-coller, qui fonctionne dans tous les cas : dans l'éditeur, "
  "menu Insertion ▸ Module ; ouvrez VP_Candidatures_a_coller.txt avec le Bloc-notes ; "
  "Ctrl + A puis Ctrl + C ; revenez dans la fenêtre blanche de l'éditeur et faites Ctrl + V.")
L("5.  Lancer l'installation",
  "Toujours dans l'éditeur : menu Exécution ▸ Exécuter Sub/UserForm, choisissez « Installer » ▸ Exécuter. "
  "Les boutons apparaissent alors dans l'onglet Accueil et en haut de l'onglet Candidatures.")
L("6.  Enregistrer", "Revenez dans Excel (Alt + F11) et enregistrez. C'est terminé.")
N("Si un bandeau jaune « Avertissement de sécurité » s'affiche à l'ouverture, cliquez sur « Activer le contenu ». "
  "Le classeur est alors considéré comme approuvé sur ce poste.")
N("Aucune référence VBA supplémentaire n'est nécessaire : le module dialogue avec Outlook et Word en liaison tardive.")

S("4.  CE QUE FAIT CHAQUE BOUTON")
BOUTONS = [
 ("Nouvelle candidature", "Ajoute une ligne, attribue l'identifiant (CAND-année-numéro) et la date du jour."),
 ("Importer depuis Outlook", "Parcourt le dossier Outlook indiqué dans Paramètres, crée une ligne par message non lu, "
                             "enregistre les pièces jointes dans le dossier du candidat et marque le message comme lu."),
 ("Accusé de réception", "Prépare l'e-mail d'accusé de réception, passe le statut à « 2-Accusé de réception » "
                          "et met à jour la date de dernier contact."),
 ("Demander des pièces", "Envoie la demande de pièces complémentaires et programme une relance."),
 ("Planifier un entretien", "Demande la date, l'heure et les modalités ; crée le rendez-vous Outlook avec invitation "
                             "au candidat, ajoute la ligne dans l'onglet Entretiens et envoie la convocation."),
 ("Mettre en vivier", "Envoie le message de mise en vivier et bascule le statut correspondant."),
 ("Réponse négative", "Envoie la réponse négative après confirmation, enregistre le motif et clôture le dossier."),
 ("Proposition / suite favorable", "Envoie le message de suite favorable et passe le statut à « 7-Proposition envoyée »."),
 ("Envoi groupé", "Applique le même modèle d'e-mail à toutes les lignes sélectionnées (utile quand un poste est pourvu)."),
 ("Fiche candidat (Word)", "Génère la fiche de synthèse et la grille d'entretien dans le dossier du candidat."),
 ("Courrier de convocation", "Génère la convocation en courrier Word, prête à être signée."),
 ("Lettre de refus", "Génère la lettre de refus sur papier à en-tête."),
 ("Convention / promesse", "Génère la convention de stage ou la promesse d'embauche à partir du modèle."),
 ("Dossier du candidat", "Crée le dossier du candidat s'il n'existe pas et l'ouvre dans l'Explorateur."),
 ("Relances à faire", "Liste les dossiers en retard et crée une tâche Outlook pour chacun."),
 ("Purge RGPD", "Recense les candidatures dont la durée de conservation est dépassée et propose leur anonymisation."),
 ("Actualiser les listes", "Étend les menus déroulants après l'ajout de valeurs dans l'onglet Paramètres."),
 ("Supprimer les exemples", "Supprime les lignes de démonstration (EXEMPLE-…) des onglets Candidatures et Entretiens."),
]
for b, d in BOUTONS:
    L(b, d)

S("5.  LE CIRCUIT TYPE D'UNE CANDIDATURE")
L("Réception", "Import Outlook ou saisie manuelle ▸ statut 1-Reçue.")
L("Sous 7 jours", "Accusé de réception ▸ statut 2. Au-delà, la colonne Alerte affiche « À traiter ».")
L("Examen", "Statut 3-En cours d'examen, attribution d'un responsable et d'une note sur 5.")
L("Entretien", "Planification ▸ statut 4, puis compte rendu dans l'onglet Entretiens ▸ statut 5.")
L("Décision", "Proposition (7) et recrutement (8), vivier (6), ou réponse négative (9).")
L("Après décision", "La ligne passe en gris. La date de purge RGPD est calculée automatiquement.")

S("6.  PROTECTION DES DONNÉES (RGPD)")
N("Un fichier de candidatures est un traitement de données personnelles : il doit figurer au registre du cabinet.")
L("Conservation", "Par défaut 2 ans à compter du dernier contact, conformément à la recommandation de la CNIL. "
                  "Le délai est modifiable dans Paramètres ; la colonne « Purge RGPD le » se recalcule seule.")
L("Information", "Les modèles d'e-mails comportent la mention d'information et le rappel des droits du candidat.")
L("Vivier", "La conservation au-delà du recrutement suppose l'accord du candidat : colonne « Consentement vivier ».")
L("Données sensibles", "Ne saisissez jamais dans les commentaires d'information sur la santé, l'origine, les opinions "
                        "politiques, religieuses ou syndicales, ni sur la vie privée du candidat.")
L("Accès", "Rangez le classeur dans un dossier à accès restreint et sauvegardez-le régulièrement.")

S("7.  DÉPANNAGE")
L("L'import du module échoue",
  "Vérifiez le nom exact du fichier : certains navigateurs le téléchargent en "
  "VP_Candidatures.bas.txt. Affichez les extensions dans l'Explorateur et renommez-le si besoin. "
  "En cas de doute, utilisez directement la variante copier-coller (étape 4 bis).")
L("Le module importé est illisible",
  "Le fichier a été réenregistré par un éditeur qui a modifié les fins de ligne ou l'encodage. "
  "Reprenez le fichier d'origine, ou utilisez la variante copier-coller.")
L("Les boutons ne réagissent pas", "Le fichier a été enregistré en .xlsx, ou les macros sont désactivées : reprenez la section 3.")
L("« Impossible de lire le fichier »", "Le fichier n'a pas été débloqué : clic droit ▸ Propriétés ▸ Débloquer.")
L("Outlook ne s'ouvre pas", "Outlook doit être installé et un profil configuré sur le poste. "
                             "La version web d'Outlook ne permet pas ce type d'automatisation.")
L("L'e-mail part sans signature", "Laissez la cellule « Signature e-mail » vide dans Paramètres : "
                                   "la signature Outlook par défaut est alors conservée.")
L("Le document Word est vide", "Vérifiez le chemin du dossier des modèles dans Paramètres et la présence des fichiers .docx.")
L("Une valeur est refusée", "Elle ne figure pas dans la liste de l'onglet Paramètres : ajoutez-la, puis cliquez sur "
                             "« Actualiser les listes ».")
L("Le classeur devient lent", "Supprimez les lignes vides sous le tableau et archivez les candidatures closes "
                               "de plus de deux ans (bouton Purge RGPD).")

r = 4
for kind, a, b in LIGNES:
    if kind == "S":
        r += 1
        wn.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = wn.cell(row=r, column=2, value=a)
        c.font = F(11, True, WHITE); c.fill = P(NAVY)
        c.alignment = Alignment(vertical="center", indent=1)
        wn.row_dimensions[r].height = 24
    elif kind == "L":
        c1 = wn.cell(row=r, column=2, value=a)
        c1.font = F(9.5, True, NAVY)
        c1.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        c2 = wn.cell(row=r, column=3, value=b)
        c2.font = F(9.5)
        c2.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        wn.row_dimensions[r].height = 15 + 12 * (len(b) // 105)
    elif kind == "N":
        wn.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = wn.cell(row=r, column=2, value="ℹ   " + b)
        c.font = F(9, False, "7D6608", it=True); c.fill = P("FCF3CF")
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        wn.row_dimensions[r].height = 16 + 12 * (len(b) // 130)
    r += 1
wn.sheet_view.showGridLines = False
wn.sheet_properties.tabColor = GREY
wn.page_setup.orientation = "portrait"
wn.page_setup.fitToWidth = 1
wn.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- ordre des onglets
ordre = ["Accueil", "Candidatures", "Entretiens", "Tableau de bord", "Journal",
         "Modèles e-mails", "Paramètres", "Notice"]
wb._sheets = [wb[n] for n in ordre]
wb.active = 0
wb.properties.title = "Gestion des candidatures — Cabinet Victimes & Préjudices"
wb.properties.creator = "Cabinet Victimes & Préjudices"
wb.properties.description = ("Suivi des candidatures, entretiens et automatisations Outlook / Word. "
                             "Module VBA associé : VP_Candidatures.bas")
import sys
OUT = sys.argv[1] if len(sys.argv) > 1 else "Gestion_Candidatures_VP.xlsx"
wb.save(OUT)
print("OK ->", OUT)
