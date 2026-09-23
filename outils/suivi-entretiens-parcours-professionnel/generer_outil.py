"""Génère l'outil Excel de suivi des entretiens de parcours professionnel (EPP).

Cadre : article L6315-1 du Code du travail issu de la loi n° 2025-989 du
24 octobre 2025 (JO du 25/10/2025, en vigueur depuis le 26/10/2025).

Usage : python3 generer_outil.py [chemin_sortie.xlsx]

Les formules n'utilisent que des fonctions disponibles depuis Excel 2010
(pas de MAXIFS / TEXTJOIN) pour rester compatibles avec les versions
d'Excel encore répandues en entreprise.
"""

import datetime as dt
import math
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

SORTIE = sys.argv[1] if len(sys.argv) > 1 else "Suivi_Entretiens_Parcours_Professionnel.xlsx"

# --------------------------------------------------------------------------
# Charte
# --------------------------------------------------------------------------
POLICE = "Arial"
MARINE = "1F3864"
BLEU_SECTION = "2F5496"
JAUNE_SAISIE = "FFF2CC"
JAUNE_ENTETE = "FFE699"
GRIS_CALCUL = "F2F2F2"
GRIS_ENTETE = "D9D9D9"
BORDURE = "BFBFBF"

STATUTS = {
    "rouge": (["En retard", "À proposer (reprise passée)", "Non réalisé", "Risque abondement"],
              "FFC7CE", "9C0006"),
    "orange": (["À planifier", "À prévoir", "À prévoir à la reprise", "À programmer avant 60 ans",
                "À vérifier (formation ?)", "À compléter"], "FCE4D6", "C65911"),
    "bleu": (["Planifié", "Planifié dans la fenêtre", "Prochain EPP dans la fenêtre"],
             "DDEBF7", "1F4E79"),
    "vert": (["À jour", "Réalisé", "Dispensé (EPP récent)", "Pas de risque"], "C6EFCE", "006100"),
    "gris": (["Sorti", "Non concerné"], "EDEDED", "7F7F7F"),
}

FMT_DATE = "DD/MM/YYYY"
PREMIERE = 6            # première ligne de données (Suivi et Journal)
NB_SALARIES = 500
NB_ENTRETIENS = 1000
DERNIERE = PREMIERE + NB_SALARIES - 1
DERNIERE_J = PREMIERE + NB_ENTRETIENS - 1

fin = Side(style="thin", color=BORDURE)
BORD = Border(left=fin, right=fin, top=fin, bottom=fin)


def police(**kw):
    kw.setdefault("name", POLICE)
    kw.setdefault("size", 10)
    return Font(**kw)


def remplir(couleur):
    return PatternFill("solid", start_color=couleur, end_color=couleur)


def titre(ws, cellule, texte, taille=14):
    ws[cellule] = texte
    ws[cellule].font = police(size=taille, bold=True, color=MARINE)


def bandeau(ws, plage, texte, couleur=MARINE):
    ws.merge_cells(plage)
    c = ws[plage.split(":")[0]]
    c.value = texte
    c.font = police(bold=True, color="FFFFFF")
    c.fill = remplir(couleur)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)


def regles_statut(ws, plage, ref):
    """Mise en forme conditionnelle des libellés de statut (ref = 1re cellule, relative)."""
    for libelles, fond, texte in STATUTS.values():
        tests = [f'{ref}="{lib}"' for lib in libelles]
        if "Non concerné" in libelles:
            tests.append(f'LEFT({ref},12)="Non concerné"')
        ws.conditional_formatting.add(
            plage,
            FormulaRule(formula=[f"OR({','.join(tests)})"], fill=remplir(fond),
                        font=Font(name=POLICE, color=texte, bold=True)))


def hauteur(texte, largeur_car):
    lignes = sum(max(1, math.ceil(len(p) / largeur_car)) for p in str(texte).split("\n"))
    return 14 * lignes + 4


wb = Workbook()
ws_guide = wb.active
ws_guide.title = "Mode d'emploi"
ws_tdb = wb.create_sheet("Tableau de bord")
ws = wb.create_sheet("Suivi EPP")
ws_j = wb.create_sheet("Journal")
ws_p = wb.create_sheet("Paramètres")

# --------------------------------------------------------------------------
# Paramètres
# --------------------------------------------------------------------------
titre(ws_p, "B1", "Paramètres de l'outil")
ws_p["B2"] = ("Cellules jaunes modifiables. Les délais reprennent le Code du travail ; "
              "adaptez-les si un accord d'entreprise ou de branche prévoit des règles plus favorables.")
ws_p["B2"].font = police(italic=True, color="595959")

for col, texte in zip("BCDE", ["Paramètre", "Valeur", "Unité", "Commentaire / base légale"]):
    c = ws_p[f"{col}4"]
    c.value = texte
    c.font = police(bold=True, color="FFFFFF")
    c.fill = remplir(MARINE)
    c.border = BORD
    c.alignment = Alignment(vertical="center")

PARAMS = [
    ("DateRef", "Date de référence du suivi", "=TODAY()", "date",
     "Par défaut : date du jour (formule =AUJOURDHUI()). Saisir une date fixe pour simuler la situation à une autre date."),
    ("DateVigueur", "Entrée en vigueur de l'entretien de parcours professionnel", dt.date(2025, 10, 26), "date",
     "Loi n° 2025-989 du 24/10/2025, publiée au JO du 25/10/2025, applicable le lendemain."),
    ("ReglesTransitoires", "Appliquer les règles transitoires", "Oui", "Oui / Non",
     "Oui : si l'ancien délai (2 ans pour l'entretien, 6 ans pour le bilan) était déjà expiré au 26/10/2025, "
     "l'échéance reste l'ancienne (retard à régulariser). Sinon, le nouveau délai court depuis le dernier entretien."),
    ("DelaiPremierEPP", "Délai du 1er EPP après l'embauche", 12, "mois",
     "Art. L6315-1 : EPP « au cours de la première année suivant son embauche »."),
    ("PeriodiciteEPP", "Périodicité de l'EPP", 4, "ans",
     "Art. L6315-1 : tous les quatre ans, à compter du dernier entretien réalisé."),
    ("AnciennePeriodiciteEP", "Ancienne périodicité de l'entretien professionnel", 2, "ans",
     "Régime antérieur au 26/10/2025 – sert uniquement aux règles transitoires."),
    ("PeriodiciteEDL", "Périodicité de l'état des lieux récapitulatif", 8, "ans",
     "Art. L6315-1 : tous les huit ans (le premier peut intervenir 7 ans après le 1er EPP)."),
    ("AnciennePeriodiciteEDL", "Ancienne périodicité du bilan récapitulatif", 6, "ans",
     "Régime antérieur – sert uniquement aux règles transitoires."),
    ("FenetreRetour", "Dispense d'EPP au retour d'absence", 12, "mois",
     "Art. L6315-1 : l'EPP de reprise n'est pas dû si un EPP a eu lieu dans les 12 mois précédant la reprise."),
    ("DelaiMiCarriere", "Délai de l'EPP après la visite de mi-carrière", 2, "mois",
     "Art. L6315-1 : EPP organisé dans les 2 mois suivant la visite médicale de mi-carrière (art. L4624-2-2)."),
    ("AgeMiCarriere", "Âge de la visite médicale de mi-carrière", 45, "ans",
     "Art. L4624-2-2 : à défaut d'accord de branche, année civile des 45 ans. Déclenche l'alerte « Visite de mi-carrière à vérifier »."),
    ("AgeFinCarriere", "Âge de référence – EPP de fin de carrière", 60, "ans",
     "Art. L6315-1 : le 1er EPP intervenant dans les 2 années précédant le 60e anniversaire aborde le maintien "
     "dans l'emploi et les aménagements de fin de carrière (temps partiel, retraite progressive)."),
    ("FenetreFinCarriere", "Fenêtre avant cet âge", 2, "ans", "Début de la fenêtre = 60e anniversaire – 2 ans."),
    ("SeuilAlerteEPP", "Seuil d'alerte EPP", 90, "jours",
     "Un EPP dont l'échéance est à moins de N jours passe au statut « À planifier »."),
    ("SeuilAlerteEDL", "Seuil d'alerte état des lieux", 180, "jours",
     "Un état des lieux dont l'échéance est à moins de N jours passe au statut « À prévoir »."),
    ("Effectif", "Effectif de l'entreprise", 60, "salariés",
     "À RENSEIGNER (60 = valeur d'exemple). Détermine l'exposition à l'abondement correctif du CPF."),
    ("SeuilEffectif", "Seuil d'effectif – abondement correctif", 50, "salariés",
     "Art. L6323-13 : entreprises d'au moins 50 salariés."),
    ("MontantAbondement", "Montant de l'abondement correctif CPF", 3000, "€",
     "Art. R6323-3 : 3 000 € par salarié concerné."),
]

for i, (nom, libelle, valeur, unite, commentaire) in enumerate(PARAMS):
    r = 5 + i
    ws_p[f"B{r}"] = libelle
    ws_p[f"C{r}"] = valeur
    ws_p[f"D{r}"] = unite
    ws_p[f"E{r}"] = commentaire
    for col in "BCDE":
        ws_p[f"{col}{r}"].border = BORD
        ws_p[f"{col}{r}"].alignment = Alignment(vertical="center", wrap_text=(col == "E"))
        ws_p[f"{col}{r}"].font = police()
    ws_p[f"C{r}"].fill = remplir(JAUNE_SAISIE)
    ws_p[f"C{r}"].font = police(color="0000FF", bold=True)
    ws_p[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center")
    if unite == "date":
        ws_p[f"C{r}"].number_format = FMT_DATE
    elif unite == "€":
        ws_p[f"C{r}"].number_format = '#,##0 "€"'
    ws_p.row_dimensions[r].height = hauteur(commentaire, 70)
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Paramètres'!$C${r}")

ligne_transitoire = 5 + [p[0] for p in PARAMS].index("ReglesTransitoires")
dv_oui_non_p = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
ws_p.add_data_validation(dv_oui_non_p)
dv_oui_non_p.add(f"C{ligne_transitoire}")

LISTES = {
    "G": ("Services", "ListeServices", 20,
          ["Direction", "Ressources humaines", "Administratif et finance", "Commercial", "Production",
           "Logistique", "Informatique", "Qualité"]),
    "H": ("Types d'absence (reprise)", "ListeAbsences", 9,
          ["Congé de maternité", "Congé d'adoption", "Congé parental d'éducation",
           "Temps partiel parental (L1225-47)", "Congé de proche aidant", "Congé sabbatique",
           "Mobilité volontaire sécurisée", "Arrêt longue maladie", "Mandat syndical"]),
    "I": ("Types d'entretien (Journal)", "ListeTypes", 8,
          ["EPP – 1re année", "EPP périodique", "EPP + état des lieux récapitulatif",
           "EPP de reprise après absence", "EPP après visite de mi-carrière",
           "EPP de fin de carrière (58-60 ans)", "Entretien professionnel (ancien régime)",
           "Bilan à 6 ans (ancien régime)"]),
}
for col, (entete, nom, taille, valeurs) in LISTES.items():
    c = ws_p[f"{col}4"]
    c.value = entete
    c.font = police(bold=True, color="FFFFFF")
    c.fill = remplir(MARINE)
    c.border = BORD
    for k in range(taille):
        cell = ws_p[f"{col}{5 + k}"]
        cell.value = valeurs[k] if k < len(valeurs) else None
        cell.fill = remplir(JAUNE_SAISIE)
        cell.font = police(color="0000FF")
        cell.border = BORD
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Paramètres'!${col}$5:${col}${4 + taille}")

# Libellés utilisés par les formules : un nom par libellé pour éviter toute faute de frappe.
wb.defined_names["TypeEDL"] = DefinedName("TypeEDL", attr_text="'Paramètres'!$I$7")
wb.defined_names["TypeBilan6"] = DefinedName("TypeBilan6", attr_text="'Paramètres'!$I$12")
ws_p["G26"] = ("Les libellés des types d'entretien « EPP + état des lieux récapitulatif » et "
               "« Bilan à 6 ans (ancien régime) » sont utilisés dans les calculs : ne pas les renommer.")
ws_p["G26"].font = police(italic=True, size=9, color="C00000")

for col, largeur in {"A": 2, "B": 46, "C": 13, "D": 10, "E": 80, "F": 3, "G": 26, "H": 32, "I": 38}.items():
    ws_p.column_dimensions[col].width = largeur
ws_p.freeze_panes = "A5"

# --------------------------------------------------------------------------
# Suivi EPP
# --------------------------------------------------------------------------
JA = f"Journal!$A${PREMIERE}:$A${DERNIERE_J}"
JC = f"Journal!$C${PREMIERE}:$C${DERNIERE_J}"
JD = f"Journal!$D${PREMIERE}:$D${DERNIERE_J}"
SORTI = 'AND($I{r}<>"",$I{r}<=DateRef)'
FEN60 = "EDATE($AB{r},-FenetreFinCarriere*12)"

ALERTES = [
    ('$H{r}=""', "Date d'embauche manquante"),
    ('AND($K{r}="Plus d\'un an",$N{r}="")', "Précédent entretien à renseigner (ancienneté > 1 an)"),
    ('$O{r}="Retard antérieur à la réforme"', "Ancien délai de 2 ans expiré avant la réforme"),
    ('AND($N{r}<>"",$H{r}<>"",$N{r}<$H{r})', "Entretien antérieur à l'embauche ?"),
    ('AND($N{r}<>"",$N{r}>DateRef)', "Entretien daté dans le futur ?"),
    ('AND($Q{r}<>"",$P{r}<>"",$Q{r}>$P{r})', "Date prévue après l'échéance"),
    ('AND($Q{r}<>"",$Q{r}<DateRef,OR($N{r}="",$N{r}<$Q{r}))', "Entretien prévu passé : le consigner au Journal"),
    ('AND($G{r}<>"",$X{r}="",YEAR(DateRef)-YEAR($G{r})>=AgeMiCarriere-1,'
     'YEAR(DateRef)-YEAR($G{r})<=AgeMiCarriere+1)', "Visite de mi-carrière à vérifier"),
]
# Concaténation « | message » puis suppression du 1er séparateur (compatible Excel 2010).
F_ALERTES = ('=IF(OR($B{r}="",$S{r}="Sorti"),"",MID('
             + "&".join(f'IF({cond}," | {msg}","")' for cond, msg in ALERTES)
             + ",4,1000))")

# (colonne, en-tête, largeur, saisie?, format, formule, commentaire d'en-tête)
COLONNES = [
    ("A", "Matricule", 11, True, None, None, "Identifiant unique : il relie le salarié au Journal des entretiens."),
    ("B", "Nom", 16, True, None, None, None),
    ("C", "Prénom", 14, True, None, None, None),
    ("D", "Service", 20, True, None, None, "Liste modifiable dans l'onglet Paramètres."),
    ("E", "Poste", 24, True, None, None, None),
    ("F", "Responsable de l'entretien", 20, True, None, None,
     "L'EPP est conduit par l'employeur (supérieur hiérarchique ou représentant de la direction), "
     "sur le temps de travail."),
    ("G", "Date de naissance", 12, True, FMT_DATE, None, "Sert aux alertes mi-carrière et fin de carrière."),
    ("H", "Date d'embauche", 12, True, FMT_DATE, None, None),
    ("I", "Date de sortie", 12, True, FMT_DATE, None, "À renseigner en cas de départ : la ligne passe en « Sorti »."),
    ("J", "Ancienneté (années)", 11, False, "0.0",
     '=IF(OR($B{r}="",$H{r}=""),"",IF($H{r}>DateRef,0,YEARFRAC($H{r},DateRef,1)))', None),
    ("K", "Ancienneté", 14, False, None,
     '=IF(OR($B{r}="",$H{r}=""),"",IF(EDATE($H{r},12)<=DateRef,"Plus d\'un an","Moins d\'un an"))',
     "Plus d'un an : le prochain EPP se calcule à partir de la date du précédent entretien professionnel."),
    ("L", "Date du précédent entretien professionnel (saisie)", 16, True, FMT_DATE, None,
     "OBLIGATOIRE pour les salariés ayant plus d'un an d'ancienneté.\n"
     "Date du dernier entretien professionnel (ancien entretien biennal) ou du dernier EPP : "
     "c'est à partir de cette date que court le délai de 4 ans.\n"
     "Les entretiens consignés dans l'onglet Journal sont repris automatiquement (la date la plus récente est retenue)."),
    ("M", "Dernier entretien consigné au Journal", 15, False, FMT_DATE,
     '=IF(OR($B{r}="",$A{r}=""),"",IF(SUMPRODUCT(MAX((' + JA + '=$A{r})*' + JC + '))=0,"",'
     'SUMPRODUCT(MAX((' + JA + '=$A{r})*' + JC + '))))', None),
    ("N", "Dernier entretien retenu", 14, False, FMT_DATE,
     '=IF($B{r}="","",IF(MAX($L{r},$M{r})=0,"",MAX($L{r},$M{r})))',
     "Date la plus récente entre la saisie (colonne L) et le Journal."),
    ("O", "Règle de calcul appliquée", 26, False, None,
     '=IF(OR($B{r}="",$H{r}=""),"",IF($N{r}="",IF(EDATE($H{r},DelaiPremierEPP)>DateRef,"1er EPP (1re année)",'
     '"Aucun entretien renseigné"),IF(AND(ReglesTransitoires="Oui",EDATE($N{r},AnciennePeriodiciteEP*12)<DateVigueur),'
     '"Retard antérieur à la réforme","Périodicité "&PeriodiciteEPP&" ans")))',
     "Moins d'un an sans entretien : 1er EPP dans les 12 mois suivant l'embauche.\n"
     "Précédent entretien connu : précédent entretien + 4 ans.\n"
     "Si l'ancien délai de 2 ans était déjà expiré au 26/10/2025 : échéance = précédent + 2 ans (retard)."),
    ("P", "Échéance du prochain EPP", 14, False, FMT_DATE,
     '=IF(OR($B{r}="",$H{r}=""),"",IF($N{r}="",EDATE($H{r},DelaiPremierEPP),'
     'IF(AND(ReglesTransitoires="Oui",EDATE($N{r},AnciennePeriodiciteEP*12)<DateVigueur),'
     'EDATE($N{r},AnciennePeriodiciteEP*12),EDATE($N{r},PeriodiciteEPP*12))))', None),
    ("Q", "Date prévue (convocation)", 14, True, FMT_DATE, None,
     "Date de l'entretien programmé. Le statut passe à « Planifié » si elle est postérieure à la date "
     "de référence et antérieure à l'échéance."),
    ("R", "Jours restants", 10, False, "0;[Red]-0",
     '=IF(OR($P{r}="",' + SORTI + '),"",$P{r}-DateRef)', None),
    ("S", "Statut EPP", 14, False, None,
     '=IF($B{r}="","",IF(' + SORTI + ',"Sorti",IF($H{r}="","À compléter",'
     'IF(AND($Q{r}<>"",$Q{r}>=DateRef,$Q{r}<=$P{r}),"Planifié",IF($R{r}<0,"En retard",'
     'IF($R{r}<=SeuilAlerteEPP,"À planifier","À jour"))))))', None),
    ("T", "Alertes / contrôles", 48, False, None, F_ALERTES, None),
    ("U", "Type d'absence", 28, True, None, None,
     "Congé de maternité, d'adoption, parental d'éducation, temps partiel parental, proche aidant, "
     "sabbatique, mobilité volontaire sécurisée, arrêt longue maladie, mandat syndical."),
    ("V", "Date de reprise", 12, True, FMT_DATE, None,
     "Date de reprise effective ou prévue. L'EPP est proposé à la reprise, sauf si un EPP a eu lieu "
     "dans les 12 mois qui la précèdent."),
    ("W", "EPP de reprise", 26, False, None,
     '=IF(OR($B{r}="",$V{r}=""),"",IF(' + SORTI + ',"",IF(AND($N{r}<>"",$N{r}>=$V{r}),"Réalisé",'
     'IF(AND($N{r}<>"",$N{r}>=EDATE($V{r},-FenetreRetour)),"Dispensé (EPP récent)",'
     'IF($V{r}>DateRef,"À prévoir à la reprise","À proposer (reprise passée)")))))', None),
    ("X", "Date de la visite de mi-carrière", 14, True, FMT_DATE, None,
     "Visite médicale de mi-carrière (service de prévention et de santé au travail). "
     "Un EPP doit être organisé dans les 2 mois qui suivent."),
    ("Y", "Échéance EPP post-visite", 14, False, FMT_DATE,
     '=IF(OR($B{r}="",$X{r}=""),"",EDATE($X{r},DelaiMiCarriere))', None),
    ("Z", "Statut EPP mi-carrière", 16, False, None,
     '=IF(OR($B{r}="",$X{r}=""),"",IF(' + SORTI + ',"",IF(AND($N{r}<>"",$N{r}>=$X{r}),"Réalisé",'
     'IF($Y{r}<DateRef,"En retard","À planifier"))))', None),
    ("AA", "Âge", 7, False, "0",
     '=IF(OR($B{r}="",$G{r}=""),"",IF($G{r}>DateRef,"",DATEDIF($G{r},DateRef,"y")))', None),
    ("AB", "Date des 60 ans", 12, False, FMT_DATE,
     '=IF(OR($B{r}="",$G{r}=""),"",EDATE($G{r},AgeFinCarriere*12))', None),
    ("AC", "EPP de fin de carrière", 28, False, None,
     '=IF(OR($B{r}="",$G{r}="",$H{r}=""),"",IF(' + SORTI + ',"",IF($AB{r}<DateVigueur,"Non concerné",'
     'IF(AND($N{r}<>"",$N{r}>=MAX(' + FEN60 + ',DateVigueur),$N{r}<=$AB{r}),"Réalisé",'
     'IF(DateRef>$AB{r},"Non réalisé",'
     'IF(AND($Q{r}<>"",$Q{r}>=' + FEN60 + ',$Q{r}<=$AB{r}),"Planifié dans la fenêtre",'
     'IF(AND(MAX($P{r},DateRef)>=' + FEN60 + ',MAX($P{r},DateRef)<=$AB{r}),"Prochain EPP dans la fenêtre",'
     'IF(MAX($P{r},DateRef)>$AB{r},"À programmer avant 60 ans",""))))))))',
     "Le 1er EPP intervenant dans les 2 ans précédant les 60 ans doit aborder le maintien dans l'emploi "
     "et les aménagements de fin de carrière (temps partiel, retraite progressive).\n"
     "« À programmer avant 60 ans » : le prochain EPP tomberait après les 60 ans."),
    ("AD", "Date du dernier état des lieux / bilan (saisie)", 16, True, FMT_DATE, None,
     "Date du dernier état des lieux récapitulatif (EPP 8 ans) ou du dernier bilan à 6 ans (ancien régime). "
     "À défaut, le cycle court depuis l'embauche. Les états des lieux consignés au Journal sont repris automatiquement."),
    ("AE", "Début du cycle en cours", 13, False, FMT_DATE,
     '=IF(OR($B{r}="",$H{r}="",' + SORTI + '),"",MAX($H{r},$AD{r},IF($A{r}="",0,SUMPRODUCT(MAX((' + JA + '=$A{r})*'
     '((' + JD + '=TypeEDL)+(' + JD + '=TypeBilan6))*' + JC + ')))))', None),
    ("AF", "Échéance de l'état des lieux", 14, False, FMT_DATE,
     '=IF($AE{r}="","",IF(AND(ReglesTransitoires="Oui",EDATE($AE{r},AnciennePeriodiciteEDL*12)<DateVigueur),'
     'EDATE($AE{r},AnciennePeriodiciteEDL*12),EDATE($AE{r},PeriodiciteEDL*12)))', None),
    ("AG", "Statut état des lieux", 16, False, None,
     '=IF($AF{r}="","",IF(' + SORTI + ',"",IF($AF{r}<DateRef,"En retard",'
     'IF($AF{r}-DateRef<=SeuilAlerteEDL,"À prévoir","À jour"))))', None),
    ("AH", "Formation non obligatoire suivie sur le cycle ?", 16, True, None, None,
     "Au moins une formation autre que celles conditionnant l'exercice d'une activité ou d'une fonction "
     "(art. L6321-2) depuis le début du cycle."),
    ("AI", "Tous les EPP dus sur le cycle réalisés ?", 16, True, None, None,
     "Oui si tous les EPP dus depuis le début du cycle ont bien eu lieu."),
    ("AJ", "Risque abondement CPF", 24, False, None,
     '=IF(OR($B{r}="",$H{r}=""),"",IF(' + SORTI + ',"",IF(Effectif<SeuilEffectif,'
     '"Non concerné (< "&SeuilEffectif&" salariés)",IF($AH{r}="Oui","Pas de risque",'
     'IF(AND(OR($AI{r}="Non",$S{r}="En retard"),$AH{r}="Non"),"Risque abondement",'
     'IF(OR($AI{r}="Non",$S{r}="En retard"),"À vérifier (formation ?)",'
     'IF($AI{r}="Oui","Pas de risque","À compléter")))))))',
     "Entreprises d'au moins 50 salariés : abondement correctif de 3 000 € sur le CPF si, sur le cycle, "
     "le salarié n'a pas bénéficié des entretiens prévus ET d'au moins une formation non obligatoire."),
    ("AK", "Commentaires", 30, True, None, None, None),
]

GROUPES = [
    ("A", "I", "IDENTIFICATION DU SALARIÉ", MARINE),
    ("J", "N", "ANCIENNETÉ & PRÉCÉDENT ENTRETIEN", "2F5496"),
    ("O", "T", "PROCHAIN ENTRETIEN DE PARCOURS PROFESSIONNEL", MARINE),
    ("U", "W", "REPRISE APRÈS ABSENCE", "2F5496"),
    ("X", "Z", "MI-CARRIÈRE", MARINE),
    ("AA", "AC", "FIN DE CARRIÈRE", "2F5496"),
    ("AD", "AJ", "ÉTAT DES LIEUX RÉCAPITULATIF (8 ANS)", MARINE),
    ("AK", "AK", "SUIVI", "2F5496"),
]

ws.merge_cells("A1:O1")
titre(ws, "A1", "Suivi des entretiens de parcours professionnel (EPP)", 15)
ws.merge_cells("A2:B2")
ws["A2"] = "Date de référence :"
ws["A2"].font = police(bold=True)
ws["C2"] = "=DateRef"
ws["C2"].number_format = FMT_DATE
ws["C2"].font = police(bold=True, color=MARINE)
ws["D2"] = ("Article L6315-1 du Code du travail – loi n° 2025-989 du 24/10/2025 (en vigueur depuis le 26/10/2025) : "
            "1er EPP dans l'année suivant l'embauche, puis tous les 4 ans à compter du précédent entretien ; "
            "état des lieux récapitulatif tous les 8 ans.")
ws["D2"].font = police(italic=True, color="595959")
ws["A3"] = ("Colonnes jaunes = saisie  ·  Colonnes grises = calcul automatique (ne pas modifier)  ·  "
            "Pour tout salarié de plus d'un an, la date du précédent entretien professionnel (colonne L) est indispensable  ·  "
            "Lignes EX-… = exemples fictifs à effacer.")
ws["A3"].font = police(size=9, color="C65911", bold=True)

for debut, fin_col, texte, couleur in GROUPES:
    bandeau(ws, f"{debut}4:{fin_col}4", texte, couleur)
ws.row_dimensions[4].height = 20

for lettre, entete, largeur, saisie, fmt, formule, note in COLONNES:
    c = ws[f"{lettre}5"]
    c.value = entete
    c.font = police(bold=True, color="000000")
    c.fill = remplir(JAUNE_ENTETE if saisie else GRIS_ENTETE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORD
    if note:
        com = Comment(note, "Outil EPP")
        com.width, com.height = 320, 150
        c.comment = com
    ws.column_dimensions[lettre].width = largeur
    for r in range(PREMIERE, DERNIERE + 1):
        cell = ws[f"{lettre}{r}"]
        cell.border = BORD
        if saisie:
            cell.fill = remplir(JAUNE_SAISIE)
            cell.font = police(color="0000FF")
        else:
            cell.fill = remplir(GRIS_CALCUL)
            cell.font = police(color="C65911", size=9) if lettre == "T" else police()
            cell.value = formule.format(r=r)
        if fmt:
            cell.number_format = fmt
        if lettre in ("J", "K", "R", "S", "W", "Z", "AA", "AC", "AG", "AJ") or fmt == FMT_DATE:
            cell.alignment = Alignment(horizontal="center")
ws.row_dimensions[5].height = 54

# Validations de saisie
plage = lambda col: f"{col}{PREMIERE}:{col}{DERNIERE}"
dv_service = DataValidation(type="list", formula1="ListeServices", allow_blank=True)
dv_absence = DataValidation(type="list", formula1="ListeAbsences", allow_blank=True)
dv_oui_non = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
dv_date = DataValidation(type="date", operator="between", formula1="1", formula2="73415", allow_blank=True,
                         showErrorMessage=True, errorTitle="Date attendue",
                         error="Saisir une date au format JJ/MM/AAAA.")
for dv in (dv_service, dv_absence, dv_oui_non, dv_date):
    ws.add_data_validation(dv)
dv_service.add(plage("D"))
dv_absence.add(plage("U"))
dv_oui_non.add(plage("AH"))
dv_oui_non.add(plage("AI"))
for col in ("G", "H", "I", "L", "Q", "V", "X", "AD"):
    dv_date.add(plage(col))

# Mise en forme conditionnelle
for col in ("S", "W", "Z", "AC", "AG", "AJ"):
    regles_statut(ws, plage(col), f"{col}{PREMIERE}")
ws.conditional_formatting.add(
    f"A{PREMIERE}:AK{DERNIERE}",
    FormulaRule(formula=[f'$S{PREMIERE}="Sorti"'], font=Font(name=POLICE, color="A6A6A6", italic=True)))

ws.freeze_panes = f"D{PREMIERE}"
ws.auto_filter.ref = f"A5:AK{DERNIERE}"
ws.sheet_view.zoomScale = 90
ws.print_title_rows = "4:5"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4

# Exemples fictifs couvrant chaque cas de figure
d = dt.date
EXEMPLES = [
    dict(A="EX-001", B="MARTIN", C="Julie", D="Commercial", E="Chargée d'affaires", F="P. Durand",
         G=d(1988, 4, 12), H=d(2026, 2, 2),
         AK="Moins d'un an : 1er EPP dans l'année suivant l'embauche"),
    dict(A="EX-002", B="BERNARD", C="Thomas", D="Production", E="Technicien de maintenance", F="A. Roux",
         G=d(1981, 9, 5), H=d(2015, 6, 15), L=d(2025, 3, 10), X=d(2026, 9, 2), AD=d(2021, 6, 20),
         AH="Oui", AI="Oui", AK="Visite de mi-carrière récente : EPP sous 2 mois"),
    dict(A="EX-003", B="PETIT", C="Sophie", D="Ressources humaines", E="Gestionnaire paie", F="M. Leroy",
         G=d(1990, 11, 22), H=d(2019, 9, 1), L=d(2023, 5, 20), U="Congé parental d'éducation", V=d(2026, 10, 5),
         AH="Oui", AK="Délai de 2 ans expiré avant la réforme + reprise à venir"),
    dict(A="EX-004", B="ROBERT", C="Karim", D="Logistique", E="Cariste", F="A. Roux",
         G=d(1967, 3, 14), H=d(2008, 1, 10), L=d(2024, 12, 15), AD=d(2020, 1, 12), AH="Oui", AI="Oui",
         AK="60 ans en mars 2027 : prochain EPP à avancer"),
    dict(A="EX-005", B="RICHARD", C="Camille", D="Informatique", E="Développeuse", F="L. Garnier",
         G=d(1995, 7, 30), H=d(2021, 3, 1),
         AK="Plus d'un an sans date de précédent entretien"),
    dict(A="EX-006", B="DURAND", C="Nicolas", D="Administratif et finance", E="Comptable", F="M. Leroy",
         G=d(1979, 2, 8), H=d(2017, 4, 4), L=d(2024, 10, 12), AH="Non", AI="Non",
         AK="Bilan à 6 ans non réalisé, aucune formation non obligatoire"),
    dict(A="EX-007", B="MOREAU", C="Léa", D="Commercial", E="Assistante commerciale", F="P. Durand",
         G=d(1992, 6, 17), H=d(2025, 11, 3), Q=d(2026, 10, 15), AK="1er EPP convoqué"),
    dict(A="EX-008", B="LAURENT", C="Hugo", D="Production", E="Opérateur de production", F="A. Roux",
         G=d(1985, 1, 3), H=d(2016, 9, 20), L=d(2023, 10, 5), U="Congé de proche aidant", V=d(2026, 3, 2),
         AH="Oui", AI="Oui", AK="Dernier EPP et bilan repris du Journal"),
    dict(A="EX-009", B="SIMON", C="Chloé", D="Qualité", E="Responsable qualité", F="Direction",
         G=d(1970, 12, 25), H=d(2012, 2, 1), I=d(2026, 7, 31), L=d(2024, 2, 14), AK="Salariée sortie"),
    dict(A="EX-010", B="MICHEL", C="Antoine", D="Logistique", E="Chef d'équipe", F="A. Roux",
         G=d(1966, 10, 11), H=d(2010, 5, 3), L=d(2024, 11, 14), AD=d(2022, 5, 10), AH="Oui", AI="Oui",
         AK="EPP de fin de carrière réalisé (Journal)"),
    dict(A="EX-011", B="FOURNIER", C="Inès", D="Commercial", E="Technico-commerciale", F="P. Durand",
         G=d(1987, 5, 9), H=d(2018, 3, 12), L=d(2024, 4, 3), U="Congé de maternité", V=d(2026, 9, 1),
         AD=d(2024, 3, 15), AH="Oui", AI="Oui", AK="Reprise passée sans EPP"),
    dict(A="EX-012", B="GIRARD", C="Mehdi", D="Informatique", E="Administrateur systèmes", F="L. Garnier",
         G=d(1981, 2, 14), H=d(2020, 1, 7), L=d(2024, 1, 22),
         AK="45 ans cette année : visite de mi-carrière à organiser"),
]
for i, ex in enumerate(EXEMPLES):
    for col, val in ex.items():
        ws[f"{col}{PREMIERE + i}"] = val

# --------------------------------------------------------------------------
# Journal des entretiens
# --------------------------------------------------------------------------
ws_j.merge_cells("A1:H1")
titre(ws_j, "A1", "Journal des entretiens réalisés", 15)
ws_j["A2"] = ("Une ligne par entretien réalisé. Pour chaque matricule, la date la plus récente alimente automatiquement "
              "l'onglet « Suivi EPP » (dernier entretien retenu, prochaine échéance, état des lieux).")
ws_j["A2"].font = police(italic=True, color="595959")
ws_j["A3"] = ("Colonnes jaunes = saisie · L'EPP donne lieu à un document écrit dont une copie est remise au salarié · "
              "Lignes EX-… = exemples fictifs à effacer.")
ws_j["A3"].font = police(size=9, color="C65911", bold=True)
bandeau(ws_j, "A4:H4", "ENTRETIENS RÉALISÉS")

COL_J = [
    ("A", "Matricule", 12, True, None),
    ("B", "Salarié", 26, False, None),
    ("C", "Date de l'entretien", 14, True, FMT_DATE),
    ("D", "Type d'entretien", 36, True, None),
    ("E", "Conduit par", 20, True, None),
    ("F", "Compte rendu écrit remis ?", 14, True, None),
    ("G", "Actions / formations envisagées", 40, True, None),
    ("H", "Observations", 36, True, None),
]
F_SALARIE = ('=IF($A{r}="","",IFERROR(INDEX(\'Suivi EPP\'!$B$6:$B$505,MATCH($A{r},\'Suivi EPP\'!$A$6:$A$505,0))'
             '&" "&INDEX(\'Suivi EPP\'!$C$6:$C$505,MATCH($A{r},\'Suivi EPP\'!$A$6:$A$505,0)),"Matricule inconnu"))')
for lettre, entete, largeur, saisie, fmt in COL_J:
    c = ws_j[f"{lettre}5"]
    c.value = entete
    c.font = police(bold=True)
    c.fill = remplir(JAUNE_ENTETE if saisie else GRIS_ENTETE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORD
    ws_j.column_dimensions[lettre].width = largeur
    for r in range(PREMIERE, DERNIERE_J + 1):
        cell = ws_j[f"{lettre}{r}"]
        cell.border = BORD
        if saisie:
            cell.fill = remplir(JAUNE_SAISIE)
            cell.font = police(color="0000FF")
        else:
            cell.fill = remplir(GRIS_CALCUL)
            cell.font = police()
            cell.value = F_SALARIE.format(r=r)
        if fmt:
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="center")
ws_j.row_dimensions[5].height = 36

wb.defined_names["ListeMatricules"] = DefinedName(
    "ListeMatricules", attr_text=f"'Suivi EPP'!$A${PREMIERE}:$A${DERNIERE}")
pj = lambda col: f"{col}{PREMIERE}:{col}{DERNIERE_J}"
dv_mat = DataValidation(type="list", formula1="ListeMatricules", allow_blank=True)
dv_type = DataValidation(type="list", formula1="ListeTypes", allow_blank=True)
dv_on_j = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
dv_date_j = DataValidation(type="date", operator="between", formula1="1", formula2="73415", allow_blank=True,
                           showErrorMessage=True, errorTitle="Date attendue",
                           error="Saisir une date au format JJ/MM/AAAA.")
for dv in (dv_mat, dv_type, dv_on_j, dv_date_j):
    ws_j.add_data_validation(dv)
dv_mat.add(pj("A"))
dv_type.add(pj("D"))
dv_on_j.add(pj("F"))
dv_date_j.add(pj("C"))
ws_j.conditional_formatting.add(
    pj("F"), FormulaRule(formula=[f'F{PREMIERE}="Non"'], fill=remplir("FFC7CE"),
                         font=Font(name=POLICE, color="9C0006", bold=True)))
ws_j.conditional_formatting.add(
    pj("B"), FormulaRule(formula=[f'B{PREMIERE}="Matricule inconnu"'], font=Font(name=POLICE, color="C00000", bold=True)))
ws_j.freeze_panes = f"A{PREMIERE}"
ws_j.auto_filter.ref = f"A5:H{DERNIERE_J}"

JOURNAL_EX = [
    ("EX-008", d(2022, 9, 15), "Bilan à 6 ans (ancien régime)", "A. Roux", "Oui", "", "Bilan des 6 ans"),
    ("EX-008", d(2025, 11, 18), "EPP périodique", "A. Roux", "Oui", "Habilitation électrique B1V", ""),
    ("EX-010", d(2026, 1, 20), "EPP de fin de carrière (58-60 ans)", "A. Roux", "Oui",
     "Étude d'une retraite progressive", "Maintien dans l'emploi et temps partiel abordés"),
]
for i, (mat, date_e, type_e, par, cr, actions, obs) in enumerate(JOURNAL_EX):
    r = PREMIERE + i
    ws_j[f"A{r}"], ws_j[f"C{r}"], ws_j[f"D{r}"] = mat, date_e, type_e
    ws_j[f"E{r}"], ws_j[f"F{r}"] = par, cr
    ws_j[f"G{r}"], ws_j[f"H{r}"] = actions or None, obs or None

# --------------------------------------------------------------------------
# Tableau de bord
# --------------------------------------------------------------------------
S = lambda col: f"'Suivi EPP'!${col}${PREMIERE}:${col}${DERNIERE}"
titre(ws_tdb, "B1", "Tableau de bord – Entretiens de parcours professionnel", 15)
ws_tdb["B2"] = "Date de référence :"
ws_tdb["C2"] = "=DateRef"
ws_tdb["C2"].number_format = FMT_DATE
ws_tdb["B3"] = "Effectif déclaré (Paramètres) :"
ws_tdb["C3"] = "=Effectif"
for c in ("B2", "B3"):
    ws_tdb[c].font = police(bold=True)
for c in ("C2", "C3"):
    ws_tdb[c].font = police(bold=True, color=MARINE)
    ws_tdb[c].alignment = Alignment(horizontal="center")

statut = lambda lib: f'COUNTIF({S("S")},"{lib}")'
ACTIFS = "+".join(statut(x) for x in ("En retard", "À planifier", "Planifié", "À jour", "À compléter"))

KPI = [
    ("section", "ENTRETIENS DE PARCOURS PROFESSIONNEL"),
    ("Salariés actifs suivis", f"={ACTIFS}", "0", None),
    ("   dont ancienneté de moins d'un an", f'=COUNTIFS({S("K")},"Moins d\'un an",{S("S")},"<>Sorti")', "0", None),
    ("   dont ancienneté de plus d'un an", f'=COUNTIFS({S("K")},"Plus d\'un an",{S("S")},"<>Sorti")', "0", None),
    ("EPP en retard", f"={statut('En retard')}", "0", "rouge"),
    ('="EPP à planifier (échéance ≤ "&SeuilAlerteEPP&" jours)"', f"={statut('À planifier')}", "0", "orange"),
    ("EPP planifiés", f"={statut('Planifié')}", "0", "bleu"),
    ("EPP à jour", f"={statut('À jour')}", "0", "vert"),
    ("Taux de conformité (salariés actifs sans retard)", "=IF(C6=0,\"\",1-C9/C6)", "0%", None),
    ("blank",),
    ("section", "CONTRÔLES DE SAISIE"),
    ("Salariés de plus d'un an sans date de précédent entretien",
     f'=COUNTIFS({S("K")},"Plus d\'un an",{S("N")},"",{S("S")},"<>Sorti")', "0", "rouge"),
    ("Lignes comportant au moins une alerte", f'=SUMPRODUCT((LEN({S("T")})>0)*1)', "0", "orange"),
    ("Fiches incomplètes (date d'embauche manquante)", f"={statut('À compléter')}", "0", "orange"),
    ("blank",),
    ("section", "CAS PARTICULIERS"),
    ("Reprise d'absence : EPP à proposer (reprise passée)",
     f'=COUNTIF({S("W")},"À proposer (reprise passée)")', "0", "rouge"),
    ("Reprise d'absence : EPP à prévoir à la reprise", f'=COUNTIF({S("W")},"À prévoir à la reprise")', "0", "orange"),
    ("Mi-carrière : EPP post-visite en retard", f'=COUNTIF({S("Z")},"En retard")', "0", "rouge"),
    ("Mi-carrière : EPP post-visite à planifier", f'=COUNTIF({S("Z")},"À planifier")', "0", "orange"),
    ("Fin de carrière : EPP à programmer avant 60 ans", f'=COUNTIF({S("AC")},"À programmer avant 60 ans")', "0", "orange"),
    ("Fin de carrière : EPP non réalisé avant 60 ans", f'=COUNTIF({S("AC")},"Non réalisé")', "0", "rouge"),
    ("blank",),
    ("section", "ÉTAT DES LIEUX RÉCAPITULATIF & ABONDEMENT CPF"),
    ("États des lieux en retard", f'=COUNTIF({S("AG")},"En retard")', "0", "rouge"),
    ('="États des lieux à prévoir (≤ "&SeuilAlerteEDL&" jours)"', f'=COUNTIF({S("AG")},"À prévoir")', "0", "orange"),
    ("Salariés exposés au risque d'abondement correctif", f'=COUNTIF({S("AJ")},"Risque abondement")', "0", "rouge"),
    ("Situations à vérifier (formation non renseignée)", f'=COUNTIF({S("AJ")},"À vérifier (formation ?)")', "0", "orange"),
    ("Exposition financière potentielle", "=IF(Effectif<SeuilEffectif,0,C31*MontantAbondement)",
     '#,##0 "€"', "rouge"),
]
COULEURS_KPI = {k: (v[1], v[2]) for k, v in STATUTS.items()}
r = 5
for item in KPI:
    if item[0] == "section":
        bandeau(ws_tdb, f"B{r}:C{r}", item[1])
    elif item[0] != "blank":
        libelle, formule, fmt, couleur = item
        ws_tdb[f"B{r}"] = libelle
        ws_tdb[f"C{r}"] = formule
        ws_tdb[f"C{r}"].number_format = fmt
        ws_tdb[f"B{r}"].font = police()
        ws_tdb[f"C{r}"].font = police(bold=True, size=11)
        ws_tdb[f"C{r}"].alignment = Alignment(horizontal="center")
        for col in "BC":
            ws_tdb[f"{col}{r}"].border = BORD
        if couleur:
            fond, texte = COULEURS_KPI[couleur]
            ws_tdb.conditional_formatting.add(
                f"C{r}", FormulaRule(formula=[f"C{r}>0"], fill=remplir(fond),
                                     font=Font(name=POLICE, color=texte, bold=True, size=11)))
    r += 1
# Contrôle de cohérence des références codées en dur ci-dessus (C6, C9, C31).
assert ws_tdb["B6"].value == "Salariés actifs suivis" and ws_tdb["B9"].value == "EPP en retard"
assert ws_tdb["B31"].value == "Salariés exposés au risque d'abondement correctif"

# Répartition par service
bandeau(ws_tdb, "E5:K5", "RÉPARTITION PAR SERVICE")
for col, texte in zip("EFGHIJK", ["Service", "Actifs", "En retard", "À planifier", "Planifiés", "À jour",
                                  "Conformité"]):
    c = ws_tdb[f"{col}6"]
    c.value = texte
    c.font = police(bold=True)
    c.fill = remplir(GRIS_ENTETE)
    c.border = BORD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for k in range(20):
    rr = 7 + k
    ws_tdb[f"E{rr}"] = f"=IF('Paramètres'!G{5 + k}=\"\",\"\",'Paramètres'!G{5 + k})"
    ws_tdb[f"F{rr}"] = (f'=IF($E{rr}="","",COUNTIFS({S("D")},$E{rr})-COUNTIFS({S("D")},$E{rr},{S("S")},"Sorti"))')
    for col, lib in zip("GHIJ", ["En retard", "À planifier", "Planifié", "À jour"]):
        ws_tdb[f"{col}{rr}"] = f'=IF($E{rr}="","",COUNTIFS({S("D")},$E{rr},{S("S")},"{lib}"))'
    ws_tdb[f"K{rr}"] = f'=IF(OR($E{rr}="",N($F{rr})=0),"",1-$G{rr}/$F{rr})'
    ws_tdb[f"K{rr}"].number_format = "0%"
    for col in "EFGHIJK":
        ws_tdb[f"{col}{rr}"].border = BORD
        ws_tdb[f"{col}{rr}"].font = police()
        if col != "E":
            ws_tdb[f"{col}{rr}"].alignment = Alignment(horizontal="center")
rt = 27
ws_tdb[f"E{rt}"] = "Total"
for col in "FGHIJ":
    ws_tdb[f"{col}{rt}"] = f"=SUM({col}7:{col}26)"
ws_tdb[f"K{rt}"] = f'=IF(F{rt}=0,"",1-G{rt}/F{rt})'
ws_tdb[f"K{rt}"].number_format = "0%"
ws_tdb[f"E{rt + 1}"] = '=IF(F27<>C6,"Écart avec le total : "&(C6-F27)&" salarié(s) sans service reconnu","")'
ws_tdb[f"E{rt + 1}"].font = police(size=9, italic=True, color="C00000")
for col in "EFGHIJK":
    ws_tdb[f"{col}{rt}"].font = police(bold=True)
    ws_tdb[f"{col}{rt}"].fill = remplir(GRIS_CALCUL)
    ws_tdb[f"{col}{rt}"].border = BORD
    if col != "E":
        ws_tdb[f"{col}{rt}"].alignment = Alignment(horizontal="center")
ws_tdb.conditional_formatting.add(
    "G7:G27", FormulaRule(formula=["N(G7)>0"], font=Font(name=POLICE, color="9C0006", bold=True)))

# Plan de charge sur 12 mois
bandeau(ws_tdb, "E30:J30", "PLAN DE CHARGE – 12 PROCHAINS MOIS")
for col, texte in zip("EFGHIJ", ["Mois", "Du", "Au", "EPP à échéance", "EPP planifiés", "États des lieux"]):
    c = ws_tdb[f"{col}31"]
    c.value = texte
    c.font = police(bold=True)
    c.fill = remplir(GRIS_ENTETE)
    c.border = BORD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws_tdb["E32"] = "Retard à résorber"
ws_tdb["H32"] = f"={statut('En retard')}"
ws_tdb["J32"] = f'=COUNTIF({S("AG")},"En retard")'
for col in "EFGHIJ":
    ws_tdb[f"{col}32"].font = police(bold=True, color="9C0006")
    ws_tdb[f"{col}32"].fill = remplir("FFC7CE")
    ws_tdb[f"{col}32"].border = BORD
    ws_tdb[f"{col}32"].alignment = Alignment(horizontal="center")
for m in range(12):
    rr = 33 + m
    ws_tdb[f"F{rr}"] = "=DateRef" if m == 0 else f"=EOMONTH(DateRef,{m - 1})+1"
    ws_tdb[f"G{rr}"] = f"=EOMONTH(DateRef,{m})"
    ws_tdb[f"E{rr}"] = f"=F{rr}"
    ws_tdb[f"E{rr}"].number_format = "MMMM YYYY"
    ws_tdb[f"F{rr}"].number_format = FMT_DATE
    ws_tdb[f"G{rr}"].number_format = FMT_DATE
    ws_tdb[f"H{rr}"] = f'=COUNTIFS({S("P")},">="&$F{rr},{S("P")},"<="&$G{rr},{S("S")},"<>Sorti")'
    ws_tdb[f"I{rr}"] = f'=COUNTIFS({S("Q")},">="&$F{rr},{S("Q")},"<="&$G{rr},{S("S")},"<>Sorti")'
    ws_tdb[f"J{rr}"] = f'=COUNTIFS({S("AF")},">="&$F{rr},{S("AF")},"<="&$G{rr},{S("S")},"<>Sorti")'
    for col in "EFGHIJ":
        ws_tdb[f"{col}{rr}"].border = BORD
        ws_tdb[f"{col}{rr}"].font = police()
        ws_tdb[f"{col}{rr}"].alignment = Alignment(horizontal="left" if col == "E" else "center")
ws_tdb["E45"] = "Total sur 12 mois"
for col in "HIJ":
    ws_tdb[f"{col}45"] = f"=SUM({col}33:{col}44)"
for col in "EFGHIJ":
    ws_tdb[f"{col}45"].font = police(bold=True)
    ws_tdb[f"{col}45"].fill = remplir(GRIS_CALCUL)
    ws_tdb[f"{col}45"].border = BORD
    ws_tdb[f"{col}45"].alignment = Alignment(horizontal="center")
ws_tdb.conditional_formatting.add("H33:H44", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                         color="5B9BD5"))
ws_tdb.conditional_formatting.add("J33:J44", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                         color="A9D08E"))

for col, largeur in {"A": 2, "B": 52, "C": 14, "D": 3, "E": 26, "F": 12, "G": 12, "H": 14, "I": 14,
                     "J": 16, "K": 12}.items():
    ws_tdb.column_dimensions[col].width = largeur
ws_tdb.sheet_view.showGridLines = False
ws_tdb.page_setup.orientation = "landscape"
ws_tdb.page_setup.fitToWidth = 1
ws_tdb.page_setup.fitToHeight = 1
ws_tdb.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

# --------------------------------------------------------------------------
# Mode d'emploi
# --------------------------------------------------------------------------
g = ws_guide
g.sheet_view.showGridLines = False
for col, largeur in {"A": 2, "B": 34, "C": 72, "D": 34}.items():
    g.column_dimensions[col].width = largeur
LARGEUR_CAR = 150

g.merge_cells("B1:D1")
titre(g, "B1", "Outil de suivi des entretiens de parcours professionnel (EPP)", 16)
g.merge_cells("B2:D2")
g["B2"] = ("Article L6315-1 du Code du travail, dans sa rédaction issue de la loi n° 2025-989 du 24 octobre 2025 "
           "(JO du 25/10/2025), applicable depuis le 26 octobre 2025.")
g["B2"].font = police(italic=True, color="595959")

ligne = [4]


def paragraphe(texte, gras=False, couleur="000000"):
    r = ligne[0]
    g.merge_cells(f"B{r}:D{r}")
    g[f"B{r}"] = texte
    g[f"B{r}"].font = police(bold=gras, color=couleur)
    g[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    g.row_dimensions[r].height = hauteur(texte, LARGEUR_CAR)
    ligne[0] += 1


def section(texte):
    r = ligne[0]
    bandeau(g, f"B{r}:D{r}", texte)
    g.row_dimensions[r].height = 20
    ligne[0] += 1


def tableau(entetes, lignes, statut_col=None):
    r = ligne[0]
    for col, texte in zip("BCD", entetes):
        c = g[f"{col}{r}"]
        c.value = texte
        c.font = police(bold=True)
        c.fill = remplir(GRIS_ENTETE)
        c.border = BORD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    r += 1
    largeurs = {"B": 34, "C": 72, "D": 34}
    for valeurs in lignes:
        h = 18
        for col, texte in zip("BCD", valeurs):
            c = g[f"{col}{r}"]
            c.value = texte
            c.font = police()
            c.border = BORD
            c.alignment = Alignment(vertical="top", wrap_text=True)
            h = max(h, hauteur(texte, largeurs[col] * 1.05))
        if statut_col:
            for libelles, fond, texte_c in STATUTS.values():
                if valeurs[0].split(" / ")[0] in libelles:
                    g[f"B{r}"].fill = remplir(fond)
                    g[f"B{r}"].font = police(bold=True, color=texte_c)
        g.row_dimensions[r].height = h
        r += 1
    ligne[0] = r + 1


section("1. MODE D'EMPLOI")
for etape in [
    "① Paramètres – Renseignez l'effectif de l'entreprise et complétez la liste des services. Vérifiez les délais : "
    "ils reprennent le Code du travail et peuvent être adaptés si un accord d'entreprise ou de branche prévoit des règles "
    "plus favorables.",
    "② Suivi EPP – Une ligne par salarié ; saisissez uniquement dans les colonnes jaunes. Pour tout salarié ayant plus "
    "d'un an d'ancienneté, renseignez la date du précédent entretien professionnel (colonne L : ancien entretien "
    "biennal ou EPP) : c'est elle qui fait courir le délai de 4 ans. Sans cette date, l'outil considère qu'aucun "
    "entretien n'a eu lieu, affiche un retard et une alerte.",
    "③ Journal – Consignez chaque nouvel entretien réalisé (matricule, date, type). La date la plus récente entre la "
    "colonne L et le Journal est retenue automatiquement : inutile de ressaisir la date dans le Suivi. Un entretien de "
    "type « EPP + état des lieux récapitulatif » ou « Bilan à 6 ans » relance aussi le cycle de 8 ans.",
    "④ Tableau de bord – Vue d'ensemble des retards, contrôles de saisie, cas particuliers, exposition à "
    "l'abondement correctif, répartition par service et plan de charge sur 12 mois.",
    "⑤ Les lignes EX-001 à EX-012 (Suivi) et les lignes EX-… du Journal sont des exemples fictifs illustrant chaque "
    "cas : effacez le contenu de leurs cellules jaunes (touche Suppr) sans supprimer les lignes, afin de conserver les "
    "formules. L'outil est préformaté pour 500 salariés et 1 000 entretiens.",
    "La date de référence (Paramètres) est par défaut la date du jour : tous les statuts se mettent à jour à chaque "
    "ouverture du fichier.",
]:
    paragraphe(etape)
ligne[0] += 1

section("2. RÈGLES DE CALCUL APPLIQUÉES")
tableau(["Situation", "Règle appliquée par l'outil", "Référence"], [
    ("Salarié de moins d'un an, sans entretien",
     "1er EPP au plus tard 12 mois après la date d'embauche.",
     "L6315-1 : EPP « au cours de la première année suivant son embauche »"),
    ("Salarié de plus d'un an (ou 1er EPP déjà réalisé)",
     "Prochain EPP = date du précédent entretien professionnel (ou du dernier EPP) + 4 ans.",
     "L6315-1 ; le délai de 4 ans court à compter du dernier entretien réalisé"),
    ("Précédent entretien dont le délai de 2 ans était déjà expiré au 26/10/2025",
     "L'échéance reste fixée à précédent entretien + 2 ans : le salarié apparaît en retard et l'EPP est à organiser "
     "sans attendre. Règle désactivable dans les Paramètres.",
     "Règle transitoire (délais non expirés prolongés)"),
    ("Salarié de plus d'un an sans date de précédent entretien",
     "Échéance = embauche + 12 mois (donc dépassée) + alerte « Précédent entretien à renseigner ».",
     "Contrôle de saisie"),
    ("Reprise après une absence",
     "EPP proposé à la reprise (congé de maternité, d'adoption, parental d'éducation ou temps partiel parental, "
     "proche aidant, sabbatique, mobilité volontaire sécurisée, arrêt longue maladie, mandat syndical), sauf si un EPP "
     "a eu lieu dans les 12 mois précédant la reprise (« Dispensé »).",
     "L6315-1"),
    ("Visite médicale de mi-carrière",
     "EPP à organiser dans les 2 mois suivant la visite. Alerte si le salarié a 44 à 46 ans (année civile) sans visite "
     "renseignée.",
     "L6315-1 ; L4624-2-2"),
    ("Approche des 60 ans",
     "Le 1er EPP intervenant dans les 2 ans précédant le 60e anniversaire aborde le maintien dans l'emploi et les "
     "aménagements de fin de carrière (temps partiel, retraite progressive). L'outil signale si le prochain EPP tomberait "
     "après les 60 ans (« À programmer avant 60 ans »).",
     "L6315-1"),
    ("État des lieux récapitulatif",
     "Tous les 8 ans à compter de l'embauche ou du dernier état des lieux / bilan. Si le bilan à 6 ans était déjà échu "
     "au 26/10/2025, l'échéance reste l'ancienne (retard) ; sinon elle est reportée à 8 ans.",
     "L6315-1"),
    ("Abondement correctif du CPF",
     "Entreprises d'au moins 50 salariés : 3 000 € versés sur le CPF si, sur le cycle, le salarié n'a pas bénéficié des "
     "entretiens prévus et d'au moins une formation non obligatoire.",
     "L6323-13 ; R6323-3"),
])

section("3. STATUTS ET CODES COULEUR")
tableau(["Statut", "Signification", "Colonnes concernées"], [
    ("En retard / Non réalisé", "Échéance dépassée : entretien à organiser en priorité.",
     "Statut EPP, mi-carrière, fin de carrière, état des lieux"),
    ("À planifier / À prévoir", "Échéance dans moins de 90 jours (EPP) ou 180 jours (état des lieux).",
     "Statut EPP, mi-carrière, état des lieux, reprise"),
    ("Planifié", "Une date de convocation est fixée avant l'échéance.", "Statut EPP, fin de carrière"),
    ("À jour / Réalisé", "Aucune action requise à ce jour.", "Toutes"),
    ("Dispensé (EPP récent)", "Reprise d'absence : un EPP a eu lieu dans les 12 mois précédant la reprise.",
     "EPP de reprise"),
    ("Risque abondement", "Ni les entretiens prévus ni une formation non obligatoire sur le cycle (≥ 50 salariés).",
     "Risque abondement CPF"),
    ("Sorti / Non concerné", "Salarié parti, ou obligation non applicable.", "Toutes"),
], statut_col=True)

section("4. POINTS DE VIGILANCE ET HYPOTHÈSES")
for point in [
    "• L'EPP porte sur les perspectives d'évolution professionnelle (compétences, qualifications, formation, mobilité, "
    "reconversion, CPF, VAE, conseil en évolution professionnelle…) et jamais sur l'évaluation du travail du salarié : "
    "il reste distinct de l'entretien annuel d'évaluation. Il se tient sur le temps de travail et donne lieu à un "
    "document écrit dont une copie est remise au salarié (colonne « Compte rendu écrit remis ? » du Journal).",
    "• Accords collectifs : si un accord d'entreprise ou de branche aménage l'EPP, adaptez les paramètres. Les accords "
    "antérieurs prévoyant une périodicité supérieure à 4 ans devaient être mis en conformité au plus tard le "
    "1er octobre 2026.",
    "• Salariés embauchés avant le 26/10/2025 et n'ayant encore eu aucun entretien : par prudence, l'outil applique le "
    "délai d'un an de la nouvelle loi.",
    "• Fenêtre des 58-60 ans : interprétation prudente – si aucun EPP n'est prévu dans la fenêtre, l'outil recommande "
    "d'en programmer un.",
    "• Abondement correctif : le critère retenu (absence cumulée des entretiens prévus et de toute formation non "
    "obligatoire) suit la lecture administrative de l'article L6323-13 ; vérifiez-le au regard des textes et précisions "
    "ministérielles en vigueur.",
    "• Cet outil est une aide au pilotage : il ne remplace pas l'analyse juridique de situations particulières "
    "(transferts de contrats, suspensions multiples, accords spécifiques). Référez-vous au texte en vigueur sur "
    "Légifrance.",
]:
    paragraphe(point)

# Ordre d'ouverture : Mode d'emploi en premier
wb.active = 0
wb.save(SORTIE)
print(f"Classeur généré : {SORTIE}")
